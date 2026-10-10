import copy,io,base64
import unittest
from openpyxl import load_workbook
import app,test_app
from inventory_profiles import SCALA_PROFILE


class InventoryProfileTests(unittest.TestCase):
    setUpClass=test_app.AppTests.__dict__['setUpClass']
    tearDownClass=test_app.AppTests.__dict__['tearDownClass']
    setUp=test_app.AppTests.setUp
    tearDown=test_app.AppTests.tearDown
    request=test_app.AppTests.request

    def create(self):
        return self.request('/api/equipment/inventories','POST',{'name':'Scalas','layout':SCALA_PROFILE})[1]['id']

    def row(self,inventory_id,asset='SC-1',**values):
        return dict(inventory_id=inventory_id,location='Lobby',custom_values={'asset_id':asset,'monitor_model':'LG 55','orientation':'Vertical',**values})

    def test_custom_only_import_fields_types_and_scope(self):
        inventory=self.create();row=self.row(inventory)
        self.assertEqual(self.request('/api/equipment/import','POST',dict(version=1,inventory_id=inventory,equipment=[row]))[1],dict(added=1,skipped=0))
        self.assertEqual(self.request('/api/equipment/import','POST',dict(version=1,inventory_id=inventory,equipment=[row]))[1],dict(added=0,skipped=1))
        saved=app.equipment_inventory(inventory)[0]
        self.assertEqual(saved['custom_values'],row['custom_values'])
        self.assertEqual(app.equipment_inventory(),[])
        self.assertEqual(self.request('/api/equipment','POST',self.row(inventory,'SC-2',orientation='Upside down'))[0],400)
        self.assertEqual(self.request('/api/equipment','POST',self.row(inventory,'SC-2',orientation='landscape'))[1]['custom_values']['orientation'],'Horizontal')
        self.assertEqual(self.request('/api/equipment/inventories')[1]['inventories'][1]['layout'],SCALA_PROFILE)

    def test_custom_changes_reset_confirmations_and_stale_writes_fail(self):
        inventory=self.create();saved=self.request('/api/equipment','POST',self.row(inventory))[1];path=f"/api/equipment/{saved['id']}"
        self.assertEqual(self.request(path+'/confirm','POST',dict(saved,item_confirmed=True))[0],200)
        self.assertEqual(self.request(path+'/review','POST',dict(saved,data_checked=True))[0],200)
        updated=self.request(path,'PUT',self.row(inventory,orientation='Horizontal'))[1]
        self.assertEqual((updated['data_checked'],updated['item_confirmed']),(0,0))
        self.assertEqual(self.request(path+'/confirm','POST',dict(saved,item_confirmed=True))[0],400)

    def test_exports_include_custom_columns_and_roundtrip_json(self):
        inventory=self.create();self.request('/api/equipment','POST',self.row(inventory))
        report=self.request(f'/api/equipment/export?inventory_id={inventory}')[1]
        self.assertEqual(report['layout'],SCALA_PROFILE)
        self.assertEqual(report['equipment'][0]['custom_values']['asset_id'],'SC-1')
        csv=self.request(f'/api/equipment/export.csv?inventory_id={inventory}')[1].decode('utf-8-sig')
        self.assertIn('ID,Location,Monitor model,Orientation,Notes,Status',csv)
        raw=self.request(f'/api/equipment/export.xlsx?inventory_id={inventory}')[1]
        sheet=load_workbook(io.BytesIO(raw)).active
        self.assertEqual(sheet.cell(1,1).value,'ID');self.assertEqual(sheet.cell(2,1).value,'SC-1')
        self.assertEqual(self.request(f'/api/equipment/export.pdf?inventory_id={inventory}')[0],200)
        content=base64.b64encode(b'ID,Location,Monitor model,Orientation\nSC-1,Lobby,LG 55,Vertical\n').decode()
        cross=self.request('/api/equipment/reconcile','POST',dict(inventory_id=inventory,filename='scala.csv',content=content))[1]
        self.assertTrue(cross['complete'])

    def test_checkbox_buttons_validation_and_hidden_data_preserved(self):
        profile=copy.deepcopy(SCALA_PROFILE)
        profile['columns'].extend([dict(key='custom_online',label='Online',type='checkbox',important=True,filter='buttons',options=[]),dict(key='custom_zone',label='Zone',type='buttons',important=False,filter='dropdown',options=['A','B'])])
        inventory=self.request('/api/equipment/inventories','POST',dict(name='Custom',layout=profile))[1]['id']
        saved=self.request('/api/equipment','POST',self.row(inventory,custom_online='true',custom_zone='A'))[1]
        self.assertEqual(saved['custom_values']['custom_online'],'Yes')
        profile['columns']=profile['columns'][:-2]
        self.request(f'/api/equipment/inventories/{inventory}','PUT',dict(name='Custom',layout=profile))
        self.assertEqual(app.equipment_inventory(inventory)[0]['custom_values']['custom_zone'],'A')
        bad=copy.deepcopy(profile);bad['columns'][0]['key']='id'
        self.assertEqual(self.request(f'/api/equipment/inventories/{inventory}','PUT',dict(name='Custom',layout=bad))[0],400)

    def test_different_custom_ids_allow_same_stock_details(self):
        inventory=self.create()
        for asset in ['SC-1','SC-2']:
            row=dict(self.row(inventory,asset),brand='LG',model='55',description='Monitor',quantity=1)
            self.assertEqual(self.request('/api/equipment','POST',row)[0],201)
        self.assertEqual(len(app.equipment_inventory(inventory)),2)
        self.assertEqual(self.request('/api/equipment','POST',self.row(inventory,'sc-1'))[0],400)

    def test_custom_device_id_header_is_retained_for_mapping(self):
        content=base64.b64encode(b'Device ID,Location,Monitor\nD1,Lobby,LG55\n').decode()
        status,result=self.request('/api/spreadsheet-preview','POST',dict(filename='custom.csv',content=content,keep_headers=['Device ID']))
        self.assertEqual(status,200)
        self.assertEqual(result['headers'][0],'Device ID')

    def test_location_merge_keeps_distinct_custom_identifiers(self):
        inventory=self.create()
        for asset,location in [('SC-1','A'),('SC-2','B')]:
            self.request('/api/equipment','POST',dict(self.row(inventory,asset),location=location,brand='LG',model='55',description='Monitor'))
        self.assertEqual(self.request('/api/equipment/locations','POST',dict(inventory_id=inventory,action='merge',source='A',target='B'))[0],200)
        records=app.equipment_inventory(inventory)
        self.assertEqual(len(records),2)
        self.assertEqual({row['custom_values']['asset_id'] for row in records},{'SC-1','SC-2'})

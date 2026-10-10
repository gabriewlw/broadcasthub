import base64
import io
import unittest
from openpyxl import Workbook
from reconciliation import reconcile, original_records
import test_app
import app


def payload(raw, filename='original.csv'):
    return {'filename':filename, 'content':base64.b64encode(raw).decode()}


def item(id=1, **values):
    row = dict(id=id, brand='Sony', model='X', description='Camera', serial_number='S1', quantity=1, location='Store', notes='')
    row.update(values)
    return row


class ReconciliationTests(unittest.TestCase):
    def test_serial_matches_detect_differences_and_missing_without_changes(self):
        source = payload(b'brand,model,description,serial_number,quantity,location\nSony,X,Camera,S1,2,Old Store\nSony,Y,Camera,S2,1,Store\n')
        report = reconcile(source,[item()])
        self.assertEqual(report['summary']['different'],1)
        self.assertEqual(report['summary']['missing'],1)
        self.assertEqual({d['field'] for d in report['rows'][0]['differences']}, {'quantity','location'})
        self.assertFalse(report['complete'])

    def test_repeated_or_blank_serial_never_proves_two_source_rows(self):
        for serial in ('S1',''):
            source = payload(f'brand,model,description,serial_number,quantity,location\nSony,X,Camera,{serial},1,Store\nSony,X,Camera,{serial},1,Store\n'.encode())
            report = reconcile(source,[item(serial_number=serial)])
            self.assertFalse(report['complete'])
            self.assertGreater(report['summary']['ambiguous'],0)
            self.assertLessEqual(sum(r['saved_id'] is not None for r in report['rows']),1)

    def test_serial_added_after_import_still_matches_original_blank_serial(self):
        source=payload(b'brand,model,description,serial_number,quantity,location\nSony,X,Camera,,1,Store\n')
        report=reconcile(source,[item(serial_number='FOUND-1')])
        self.assertEqual(report['summary']['needs_review'],1)
        self.assertEqual(report['summary']['missing'],0)
        self.assertEqual(report['rows'][0]['saved_id'],1)

    def test_all_worksheets_auto_headers_and_double_locations(self):
        book=Workbook();book.active.title='Cover';book.active.append(['Original inventory'])
        sheet=book.create_sheet('Edit Room Deck 7')
        sheet.append(['Title']);sheet.append(['LOCATION','BRAND','MODEL','SERIAL NUMBER','QUANTITY','LOCATION','NOTES'])
        sheet.append(['Top shelf','Sony','X','S1',1,None,'Used'])
        tools=book.create_sheet('Tools');tools.append(['Location','Make','Model','Serial number','Quantity in Stock'])
        tools.append(['Drawer','Driver','N/A','N/A',6])
        raw=io.BytesIO();book.save(raw)
        records,errors=original_records(payload(raw.getvalue(),'original.xlsx'))
        self.assertEqual(len(records),2)
        self.assertEqual(errors[0]['sheet'],'Cover')
        self.assertEqual(records[0]['record']['location'],'Edit Room Deck 7')
        self.assertEqual(records[0]['details']['storage position'],'Top shelf')
        self.assertEqual(records[1]['record']['location'],'Tools')
        self.assertEqual(records[1]['record']['quantity'],6)
        self.assertEqual(records[1]['record']['serial_number'],'')

    def test_complete_only_with_every_row_definitely_matched(self):
        source=payload(b'brand,model,description,serial_number,quantity,location\nSony,X,Camera,S1,1,Store\n')
        report=reconcile(source,[item(),item(id=2,serial_number='Extra')])
        self.assertTrue(report['complete'])
        self.assertEqual(report['extra_records'][0]['id'],2)


class ReconciliationHTTPTests(unittest.TestCase):
    setUpClass=test_app.AppTests.__dict__['setUpClass']
    tearDownClass=test_app.AppTests.__dict__['tearDownClass']
    setUp=test_app.AppTests.setUp
    tearDown=test_app.AppTests.tearDown
    request=test_app.AppTests.request

    def test_read_only_comparison_scopes_inventory(self):
        other=app.save_equipment_inventory({'name':'Other'})['id']
        app.save_equipment(dict(item(),inventory_id=other))
        before=app.equipment_inventory(other)
        source=payload(b'brand,model,description,serial_number,quantity,location\nSony,X,Camera,S1,1,Store\n')
        self.assertEqual(self.request('/api/equipment/reconcile','POST',source)[1]['summary']['missing'],1)
        self.assertTrue(self.request('/api/equipment/reconcile','POST',dict(source,inventory_id=other))[1]['complete'])
        self.assertEqual(app.equipment_inventory(other),before)

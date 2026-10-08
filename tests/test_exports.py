import io
import csv
import re
import shutil
import subprocess
import unittest
from urllib.request import urlopen

from openpyxl import load_workbook

import app
import test_app


class ExportTests(unittest.TestCase):
    setUpClass = test_app.AppTests.__dict__['setUpClass']
    tearDownClass = test_app.AppTests.__dict__['tearDownClass']
    setUp = test_app.AppTests.setUp
    tearDown = test_app.AppTests.tearDown
    request = test_app.AppTests.request

    def download(self, path, mime, filename):
        with urlopen(self.url + path) as response:
            self.assertEqual(response.status, 200)
            self.assertEqual(response.headers['Content-Type'], mime)
            self.assertIn(filename, response.headers['Content-Disposition'])
            self.assertEqual(response.headers['Cache-Control'], 'no-store')
            return response.read()

    def xlsx(self, equipment=False):
        scope = '/api/equipment/export' if equipment else '/api/export'
        filename = 'broadcasthub-equipment.xlsx' if equipment else 'broadcasthub-network.xlsx'
        raw = self.download(scope + '.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename)
        return load_workbook(io.BytesIO(raw), data_only=False)

    def test_network_workbook_latest_fields_status_and_literal_formulas(self):
        device = self.request('/api/devices', 'POST', test_app.EXAMPLE)[1]
        path = f"/api/devices/{device['id']}"
        self.request(path + '/confirm', 'POST', dict(ip=device['ip'], vlan=device['vlan']))
        self.request(path + '/system', 'POST', dict(discipline='Audio'))
        self.request(path + '/notes', 'POST', dict(notes='=SUM(1,2)\nAçúcar <&>'))
        self.request('/api/devices', 'POST', dict(name='Dynamic', ip='DHCP'))
        self.request('/api/devices', 'POST', dict(name='Unassigned'))
        self.request('/api/devices', 'POST', dict(record_type='iptv', name='News', ip='239.1.1.5', port=5000, channel_source='Satellite', notes='Current channel notes'))
        before = app.inventory()
        book = self.xlsx()
        self.assertEqual(book.properties.creator, 'avtrack')
        self.assertTrue(book.properties.title.startswith('avtrack'))
        self.assertEqual(book.sheetnames, ['AV devices','IPTV channels'])
        sheet = book['AV devices']
        self.assertEqual([cell.value for cell in sheet[1]], ['Venue','Device name','IP Address','VLAN','System','Notes','IP confirmation','Category'])
        rows = {cells[1].value: cells for cells in list(sheet)[1:]}
        cells = rows['ATEM video switcher']
        self.assertEqual(sheet.cell(1, 1).font.name, 'JetBrains Mono')
        self.assertEqual(cells[1].font.name, 'DM Sans')
        self.assertEqual([cell.value for cell in cells], ['Liquid Lounge','ATEM video switcher','10.24.176.66',1500,'Audio','=SUM(1,2)\nAçúcar <&>','Confirmed','Video switcher'])
        self.assertEqual(cells[5].data_type, 's')
        self.assertEqual(cells[4].fill.fgColor.rgb, '00CFEADB')
        self.assertEqual(rows['Dynamic'][2].value, 'DHCP')
        self.assertIsNone(rows['Dynamic'][6].value)
        self.assertIsNone(rows['Unassigned'][2].value)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:H4')
        self.assertEqual(list(book['IPTV channels'].values)[1], ('News','239.1.1.5',5000,'Satellite','Current channel notes',None))
        self.assertEqual(app.inventory(), before)

    def test_equipment_workbook_preserves_blank_zero_and_all_fields(self):
        self.request('/api/equipment', 'POST', dict(brand='Sony', model='X', description='Spare', serial_number='SER-1', quantity=0, location='Store', notes='After inspection'))
        self.request('/api/equipment', 'POST', dict(description='Unknown spare'))
        book = self.xlsx(True)
        self.assertEqual(book.sheetnames, ['Equipment'])
        values = list(book.active.values)
        self.assertEqual(values[0], ('Item','Brand','Model','Serial number','Quantity','Location','Found','Notes'))
        self.assertIn(('Spare','Sony','X','SER-1',0,'Store','To find','After inspection'), values)
        self.assertIn(('Unknown spare',None,None,None,None,None,'To find',None), values)

    @unittest.skipUnless(shutil.which('pdftotext'), 'PDF text validation needs optional pdftotext')
    def test_equipment_reports_group_by_location_and_include_found_status(self):
        for location, item_name, serial in [('Z Store','Last item','Z-1'), ('A Theater','First item','A-1'), ('','Unassigned item','U-1')]:
            item = self.request('/api/equipment', 'POST', dict(description=item_name, brand='Sony', model='X', serial_number=serial, location=location))[1]
            if location == 'A Theater':
                self.request(f"/api/equipment/{item['id']}/confirm", 'POST', dict(item, item_confirmed=True))
        raw = self.request('/api/equipment/export.pdf', 'POST', dict(ids=[1,2,3]))[1]
        text = subprocess.run(['pdftotext','-layout','-','-'],input=raw,stdout=subprocess.PIPE,check=True).stdout.decode()
        self.assertLess(text.index('Location: Unassigned location'), text.index('Location: A Theater'))
        self.assertLess(text.index('Location: A Theater'), text.index('Location: Z Store'))
        headings = next(line for line in text.splitlines() if 'SERIAL NUMBER' in line)
        self.assertLess(headings.index('ITEM'), headings.index('BRAND'))
        self.assertLess(headings.index('MODEL'), headings.index('SERIAL NUMBER'))
        self.assertIn('Found', text)
        self.assertIn('To find', text)
        self.assertIn('by broadcastgab.com', text)
        book = self.xlsx(True)
        row = next(cells for cells in list(book.active)[1:] if cells[0].value == 'First item')
        self.assertEqual(row[6].value, 'Found')
        self.assertEqual(row[6].fill.fgColor.rgb, '00D1EAD7')

    @unittest.skipUnless(shutil.which('pdftotext'), 'PDF text validation needs optional pdftotext')
    def test_equipment_pdf_repeats_location_on_continuation_pages(self):
        rows = [dict(description=f'Inventory item {index}', serial_number=f'SN-{index}', location='Main lounge',
                     notes='Long inspection note ' * 80 if index == 0 else '', item_confirmed=index % 2)
                for index in range(45)]
        self.assertEqual(self.request('/api/equipment/import', 'POST', dict(version=1,equipment=rows))[0], 200)
        raw = self.request('/api/equipment/export.pdf')[1]
        text = subprocess.run(['pdftotext','-','-'],input=raw,stdout=subprocess.PIPE,check=True).stdout.decode()
        pages = [page for page in text.split('\f') if page.strip()]
        self.assertGreater(len(pages), 1)
        for page in pages:
            self.assertIn('Location: Main lounge', page)
            self.assertIn('SERIAL NUMBER', page)
        self.assertIn('Inventory item 44', text)

    def test_empty_reports_download_without_inventing_records(self):
        self.assertEqual(self.xlsx()['AV devices'].max_row, 1)
        self.assertEqual(self.xlsx(True).active.max_row, 1)
        for scope, filename in [('/api/export','broadcasthub-network.pdf'), ('/api/equipment/export','broadcasthub-equipment.pdf')]:
            raw = self.download(scope + '.pdf', 'application/pdf', filename)
            self.assertTrue(raw.startswith(b'%PDF-'))
            self.assertIn(b'%%EOF', raw)
        self.assertEqual(app.inventory(), [])
        self.assertEqual(app.equipment_inventory(), [])

    def test_scoped_exports_use_only_requested_records_and_latest_saved_values(self):
        camera = self.request('/api/devices', 'POST', dict(name='Selected camera', ip='10.20.0.1', venue='Theatre', discipline='Video'))[1]
        panel = self.request('/api/devices', 'POST', dict(name='Selected panel', ip='10.20.0.2', venue='Lounge', discipline='Video'))[1]
        self.request('/api/devices', 'POST', dict(name='Excluded DSP', ip='10.20.0.3', discipline='Audio'))
        channel = self.request('/api/devices', 'POST', dict(record_type='iptv', name='Selected channel', ip='239.20.0.1', port=5000))[1]
        self.request(f"/api/devices/{camera['id']}/notes", 'POST', dict(notes='Latest saved note'))
        self.request(f"/api/devices/{camera['id']}/confirm", 'POST', dict(ip=camera['ip'], vlan=camera['vlan']))
        payload = {'ids':[panel['id'], camera['id'], camera['id'], 999999]}
        status, exported = self.request('/api/export', 'POST', payload)
        self.assertEqual(status, 200)
        self.assertEqual([row['id'] for row in exported['devices']], [panel['id'], camera['id']])
        self.assertEqual(exported['devices'][1]['notes'], 'Latest saved note')
        self.assertEqual(exported['devices'][1]['ip_confirmed'], 1)
        status, raw = self.request('/api/export.csv', 'POST', payload)
        self.assertEqual(status, 200)
        self.assertEqual([row['name'] for row in csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))], ['Selected panel', 'Selected camera'])
        status, raw = self.request('/api/export.xlsx', 'POST', payload)
        self.assertEqual(status, 200)
        book = load_workbook(io.BytesIO(raw))
        self.assertEqual([row[1] for row in list(book['AV devices'].values)[1:]], ['Selected panel', 'Selected camera'])
        self.assertEqual(book['IPTV channels'].max_row, 1)
        status, raw = self.request('/api/export.pdf', 'POST', payload)
        self.assertEqual(status, 200)
        self.assertIn(b'http://10.20.0.1', raw)
        self.assertIn(b'http://10.20.0.2', raw)
        self.assertNotIn(b'http://10.20.0.3', raw)
        self.assertNotIn(b'http://239.20.0.1', raw)
        selected_channel = self.request('/api/export', 'POST', {'ids':[channel['id']]})[1]
        self.assertEqual([row['record_type'] for row in selected_channel['devices']], ['iptv'])
        self.assertEqual(len(self.request('/api/export')[1]['devices']), 4)

    def test_scoped_equipment_formats_and_empty_selection(self):
        selected = self.request('/api/equipment', 'POST', dict(brand='Selected brand', model='Panel', location='Theatre'))[1]
        self.request('/api/equipment', 'POST', dict(brand='Excluded brand', model='DSP', location='Store'))
        payload = {'ids':[selected['id']]}
        exported = self.request('/api/equipment/export', 'POST', payload)[1]
        self.assertEqual([row['id'] for row in exported['equipment']], [selected['id']])
        raw = self.request('/api/equipment/export.csv', 'POST', payload)[1]
        self.assertEqual([row['brand'] for row in csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))], ['Selected brand'])
        raw = self.request('/api/equipment/export.xlsx', 'POST', payload)[1]
        self.assertEqual([row[1] for row in list(load_workbook(io.BytesIO(raw)).active.values)[1:]], ['Selected brand'])
        raw = self.request('/api/equipment/export.pdf', 'POST', payload)[1]
        self.assertTrue(raw.startswith(b'%PDF-'))
        if shutil.which('pdftotext'):
            text = subprocess.run(['pdftotext', '-', '-'], input=raw, stdout=subprocess.PIPE, check=True).stdout.decode()
            self.assertIn('Selected brand', text)
            self.assertNotIn('Excluded brand', text)
        for base, key in [('/api/export', 'devices'), ('/api/equipment/export', 'equipment')]:
            self.assertEqual(self.request(base, 'POST', {'ids':[]})[1][key], [])
            self.assertEqual(len(list(csv.DictReader(io.StringIO(self.request(base+'.csv', 'POST', {'ids':[]})[1].decode('utf-8-sig'))))), 0)
            book = load_workbook(io.BytesIO(self.request(base+'.xlsx', 'POST', {'ids':[]})[1]))
            self.assertTrue(all(sheet.max_row == 1 for sheet in book))
            self.assertTrue(self.request(base+'.pdf', 'POST', {'ids':[]})[1].startswith(b'%PDF-'))

    def test_invalid_export_scope_never_falls_back_to_all_records(self):
        self.request('/api/devices', 'POST', test_app.EXAMPLE)
        for base in ['/api/export', '/api/equipment/export']:
            for payload in [{}, [], {'ids':None}, {'ids':'all'}, {'ids':[True]}, {'ids':['1']}, {'ids':[0]}, {'ids':[-1]}]:
                with self.subTest(base=base, payload=payload):
                    self.assertEqual(self.request(base, 'POST', payload)[0], 400)
            self.assertEqual(self.request(base, 'POST', {'ids':[]}, headers={'Content-Type':'application/json', 'Origin':'https://other.example'})[0], 403)
            self.assertEqual(self.request(base, 'POST', headers={'Content-Type':'application/json'})[0], 400)

    def test_tab_reports_include_only_the_requested_inventory_even_when_empty(self):
        av = self.request('/api/devices', 'POST', dict(name='AV only camera', ip='10.30.0.1'))[1]
        channel = self.request('/api/devices', 'POST', dict(record_type='iptv', name='IPTV only news', ip='239.30.0.1', port=5000))[1]
        self.request('/api/equipment', 'POST', dict(brand='Equipment only brand'))
        for record_type, name, excluded, sheet_name, other_sheet in [
            ('device', 'AV only camera', 'IPTV only news', 'AV devices', 'IPTV channels'),
            ('iptv', 'IPTV only news', 'AV only camera', 'IPTV channels', 'AV devices'),
        ]:
            for ids in [None, [av['id'], channel['id']], []]:
                payload = {'record_type':record_type}
                if ids is not None:
                    payload['ids'] = ids
                with self.subTest(record_type=record_type, ids=ids):
                    status, data = self.request('/api/export', 'POST', payload)
                    self.assertEqual(status, 200)
                    self.assertEqual([row['name'] for row in data['devices']], [] if ids == [] else [name])
                    raw = self.request('/api/export.csv', 'POST', payload)[1]
                    rows = list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
                    self.assertEqual([row['name'] for row in rows], [] if ids == [] else [name])
                    raw = self.request('/api/export.xlsx', 'POST', payload)[1]
                    book = load_workbook(io.BytesIO(raw))
                    self.assertEqual(book.sheetnames, [sheet_name])
                    self.assertEqual(book.active.max_row, 1 if ids == [] else 2)
                    raw = self.request('/api/export.pdf', 'POST', payload)[1]
                    self.assertTrue(raw.startswith(b'%PDF-'))
                    if shutil.which('pdftotext'):
                        text = subprocess.run(['pdftotext','-','-'], input=raw, stdout=subprocess.PIPE, check=True).stdout.decode()
                        self.assertIn(sheet_name, text)
                        self.assertNotIn(other_sheet, text)
                        self.assertNotIn(excluded, text)
                        self.assertNotIn('Equipment only brand', text)
            exported = self.request('/api/export?record_type='+record_type)[1]['devices']
            self.assertEqual([row['name'] for row in exported], [name])
            raw = self.request('/api/export.xlsx?record_type='+record_type)[1]
            self.assertEqual(load_workbook(io.BytesIO(raw)).sheetnames, [sheet_name])
        self.assertEqual(self.request('/api/export', 'POST', {'record_type':'equipment'})[0], 400)
        self.assertEqual(self.request('/api/export?record_type=equipment')[0], 400)
        self.assertEqual(self.request('/api/export', 'POST', {'record_type':'device', 'ids':None})[0], 400)

    def test_pdf_ipv4_links_include_channels_without_linking_blank_or_dhcp(self):
        self.request('/api/devices', 'POST', dict(name='Fixed', ip='10.24.176.99', vlan=1500))
        self.request('/api/devices', 'POST', dict(name='Dynamic', ip='DHCP'))
        self.request('/api/devices', 'POST', dict(name='Blank'))
        self.request('/api/devices', 'POST', dict(record_type='iptv', name='Channel', ip='239.1.1.1', port=1234))
        raw = self.download('/api/export.pdf', 'application/pdf', 'broadcasthub-network.pdf')
        self.assertIn(b'http://10.24.176.99', raw)
        self.assertIn(b'http://239.1.1.1', raw)
        self.assertNotIn(b'http://DHCP', raw)
        self.assertIn(b'/S /URI', raw)

    @unittest.skipUnless(shutil.which('pdffonts'), 'Embedded font validation needs optional pdffonts')
    def test_pdf_uses_avtrack_wordmark_and_embeds_website_fonts(self):
        device = self.request('/api/devices', 'POST', test_app.EXAMPLE)[1]
        self.request(f"/api/devices/{device['id']}/confirm", 'POST', dict(ip=device['ip'], vlan=device['vlan']))
        for scope, filename in [('/api/export', 'broadcasthub-network.pdf'),
                                ('/api/equipment/export', 'broadcasthub-equipment.pdf')]:
            raw = self.download(scope + '.pdf', 'application/pdf', filename)
            self.assertIn(b'/Author (avtrack)', raw)
            self.assertIn(b'/Title (avtrack - ', raw)
            if shutil.which('pdftotext'):
                text = subprocess.run(['pdftotext','-','-'], input=raw, stdout=subprocess.PIPE, check=True).stdout.decode()
                self.assertIn('avtrack', text)
                self.assertIn('REPORT GENERATED BY AVTRACK', text)
                self.assertNotIn('BROADCASTHUB', text)
            listing = subprocess.run(['pdffonts', '-'], input=raw, stdout=subprocess.PIPE,
                                     stderr=subprocess.PIPE, check=True).stdout.decode()
            fonts = ['DMSans-Regular', 'SpaceGrotesk-Bold', 'JetBrainsMono-Regular', 'JetBrainsMono-Medium']
            if scope == '/api/export':
                fonts.append('DMSans-Bold')
            for font in fonts:
                line = next(line for line in listing.splitlines() if font in line)
                self.assertRegex(line, r'TrueType\s+\S+\s+yes\s+yes\s+yes')
            self.assertNotIn('Courier', listing)
            self.assertNotIn('Vera', listing)

    @unittest.skipUnless(shutil.which('pdftotext'), 'PDF text validation needs optional pdftotext')
    def test_pdf_saved_information_pagination_and_long_notes(self):
        notes = 'Açúcar <&> ' + 'Long note ' * 195 + 'END_NOTE'
        rows = [dict(name=f'Camera {i}', venue='MAIN LOUNGE', discipline='Video', ip='DHCP', notes=notes if i == 0 else 'Reviewed') for i in range(75)]
        self.assertEqual(self.request('/api/import', 'POST', dict(version=1, devices=rows))[0], 200)
        device = self.request('/api/devices', 'POST', test_app.EXAMPLE)[1]
        self.request(f"/api/devices/{device['id']}/confirm", 'POST', dict(ip=device['ip'], vlan=device['vlan']))
        self.request('/api/devices', 'POST', dict(record_type='iptv', name='BBC News', ip='239.1.1.1', port=1234, channel_source='Onboard'))
        raw = self.download('/api/export.pdf', 'application/pdf', 'broadcasthub-network.pdf')
        text = subprocess.run(['pdftotext','-','-'], input=raw, stdout=subprocess.PIPE, check=True).stdout.decode()
        self.assertTrue(max(int(n) for n in re.findall(rb'/Count (\d+)', raw)) > 1)
        for value in ['Açúcar <&>', 'END_NOTE', 'Camera 74','Confirmed','1500','BBC News','Onboard','1234','REPORT GENERATED BY AVTRACK','broadcastgab.com']:
            self.assertIn(value, text)
        self.assertGreater(text.count('VENUE'), 1)
        self.assertIn('VLAN', text)
        self.assertIn('CONFIRMATION', text)
        self.assertNotIn('IP / VLAN / CONFIRMATION', text)
        self.assertIn(b'https://broadcastgab.com', raw)
        self.request('/api/equipment', 'POST', dict(brand='Neutrik', model='XLR', quantity=0, notes='Equipment notes'))
        raw = self.download('/api/equipment/export.pdf', 'application/pdf', 'broadcasthub-equipment.pdf')
        text = subprocess.run(['pdftotext','-','-'], input=raw, stdout=subprocess.PIPE, check=True).stdout.decode()
        for value in ['Neutrik','XLR','QUANTITY','Equipment notes']:
            self.assertIn(value, text)


if __name__ == '__main__':
    unittest.main()

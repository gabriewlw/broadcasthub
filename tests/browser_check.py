"""Optional end-to-end check: requires Playwright and Chromium."""
import csv
import json
import io
import shutil
import subprocess
import sys
import tempfile
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

with tempfile.TemporaryDirectory() as temp:
    app.DB_PATH = Path(temp) / 'inventory.sqlite3'
    server = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path=shutil.which("chromium") or None)
            context = browser.new_context(viewport={'width':390,'height':844}, is_mobile=True, has_touch=True, device_scale_factor=2)
            page = context.new_page()
            errors = []

            def accept_rows(count, total=None):
                total = count if total is None else total
                for index in range(count):
                    page.locator('#row-review-dialog').wait_for(state='visible')
                    progress = page.locator('#row-review-progress').inner_text()
                    assert f'{index + 1} of {total}' in progress
                    page.locator('#accept-review-row').click()
                    if index + 1 < count:
                        page.wait_for_function("previous => document.getElementById('row-review-dialog').open && document.getElementById('row-review-progress').textContent !== previous", arg=progress)
                    else:
                        page.locator('#spreadsheet-dialog').wait_for(state='hidden')

            def open_export_menu(panel='inventory', scope='all'):
                menu = page.locator('#' + panel + ' .export-dropdown')
                menu.locator('summary').click()
                assert menu.get_attribute('open') is not None
                assert menu.locator('a[download]').count() == 4
                menu.locator('.export-scope').select_option(scope)

            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto(f'http://127.0.0.1:{server.server_port}')
            navigation = page.get_by_role('navigation', name='Main navigation')
            nav_toggle = page.locator('#nav-toggle')
            assert nav_toggle.is_visible() and navigation.is_hidden()
            assert nav_toggle.get_attribute('aria-expanded') == 'false'
            nav_toggle.focus(); nav_toggle.press('Enter')
            assert navigation.is_visible()
            assert navigation.get_by_role('link', name='Devices', exact=True).is_visible()
            nav_toggle.press('Tab')
            assert navigation.get_by_role('link', name='Workspace', exact=True).evaluate('(node) => node === document.activeElement')
            page.keyboard.press('Escape')
            assert navigation.is_hidden()
            assert nav_toggle.evaluate('(node) => node === document.activeElement')
            nav_toggle.click()
            page.locator('#hero-intro').click()
            assert navigation.is_hidden()
            nav_toggle.click()
            navigation.get_by_role('link', name='Devices', exact=True).click()
            assert navigation.is_hidden() and nav_toggle.get_attribute('aria-expanded') == 'false'
            nav_toggle.click()
            page.set_viewport_size({'width':1440,'height':1000})
            assert nav_toggle.is_hidden() and navigation.is_visible()
            page.set_viewport_size({'width':390,'height':844})
            assert navigation.is_hidden() and nav_toggle.get_attribute('aria-expanded') == 'false'
            page.evaluate('window.scrollTo(0, 0)')
            menu = page.locator('#inventory .export-dropdown')
            assert menu.get_attribute('open') is None
            summary = menu.locator('summary')
            summary.focus(); summary.press('Enter')
            assert menu.locator('a').first.is_visible()
            assert menu.locator('.export-links').evaluate('(node) => { const rect = node.getBoundingClientRect(); return rect.left >= 0 && rect.right <= innerWidth; }')
            summary.press('Escape')
            assert menu.get_attribute('open') is None
            assert summary.evaluate('(node) => node === document.activeElement')
            open_export_menu()
            page.locator('#hero-title').click()
            assert menu.get_attribute('open') is None
            page.get_by_role('button', name='Fill ATEM example').click()
            page.get_by_role('button', name='Save device', exact=True).click()
            page.locator('.device-name').filter(has_text='ATEM video switcher').wait_for()
            assert app.inventory()[0]['ip'] == '10.24.176.66'
            assert page.locator('#total').inner_text() == '1'
            confirm = page.get_by_role('button', name='Confirm IP for ATEM video switcher', exact=True)
            assert 'pending' in confirm.get_attribute('class')
            confirm.click()
            green = page.get_by_role('button', name='IP confirmed for ATEM video switcher', exact=True)
            green.wait_for()
            assert 'confirmed' in green.get_attribute('class')
            assert green.is_disabled()
            page.reload()
            page.get_by_role('button', name='IP confirmed for ATEM video switcher', exact=True).wait_for()
            assert app.inventory()[0]['ip_confirmed'] == 1

            page.locator('#system-buttons').get_by_role('button', name='Audio', exact=True).click()
            assert page.locator('#no-results').is_visible()
            page.get_by_role('button', name='Clear', exact=True).click()
            page.locator('#venue-filter').select_option('Liquid Lounge')
            page.locator('#system-filter').select_option('Video')
            assert page.locator('#system-buttons').get_by_role('button', name='Video', exact=True).get_attribute('aria-pressed') == 'true'
            assert page.locator('.device-row').count() == 1
            page.locator('#search').fill('not found')
            assert page.locator('#no-results').is_visible()
            page.locator('#search').fill('10.24.176.66')
            assert page.locator('.device-row').count() == 1
            page.get_by_role('button', name='Edit ATEM video switcher', exact=True).click()
            page.locator('[name=ip]').fill('10.24.176.999')
            page.get_by_role('button', name='Save changes').click()
            page.locator('#form-error').wait_for(state='visible')
            assert 'valid IPv4' in page.locator('#form-error').inner_text()
            page.locator('[name=ip]').fill('10.24.176.66')
            page.locator('[name=name]').fill('ATEM main')
            page.get_by_role('button', name='Save changes').click()
            page.get_by_role('button', name='Edit ATEM main').wait_for()
            assert page.get_by_role('button', name='IP confirmed for ATEM main', exact=True).is_visible()
            open_export_menu('inventory')
            with page.expect_download() as download:
                page.get_by_role('link',name='JSON ↓').click()
            assert menu.get_attribute('open') is None
            payload = json.loads(Path(download.value.path()).read_text())
            assert payload['devices'][0]['name'] == 'ATEM main'
            page.get_by_role('button',name='Delete ATEM main').click()
            page.get_by_role('button',name='Delete device',exact=True).click()
            page.locator('#empty').wait_for(state='visible')
            path = Path(temp) / 'inventory.json'
            path.write_text(json.dumps(payload))
            page.locator('#import-file').set_input_files(str(path))
            page.get_by_role('button',name='Edit ATEM main').wait_for()
            page.reload()
            page.get_by_role('button',name='Edit ATEM main').wait_for()
            # A second venue verifies combined button filters and form choices.
            page.get_by_role('button', name='Add device', exact=True).click()
            page.locator('[name=name]').fill('Audio console')
            page.locator('[name=category]').fill('Audio console')
            page.locator('#form-system').select_option('Audio')
            page.locator('#form-venue-buttons').get_by_role('button', name='Liquid Lounge', exact=True).click()
            assert page.locator('[name=venue]').input_value() == 'Liquid Lounge'
            page.locator('[name=venue]').fill('Theater')
            assert page.locator('#form-venue-buttons button').get_attribute('aria-pressed') == 'false'
            page.locator('[name=ip]').fill('10.24.176.67')
            page.locator('[name=vlan]').fill('1500')
            page.get_by_role('button', name='Save device', exact=True).click()
            theater = page.locator('#venue-filter')
            theater.select_option('Theater')
            page.locator('#system-buttons').get_by_role('button', name='Audio', exact=True).click()
            assert theater.input_value() == 'Theater'
            assert page.locator('.device-row').count() == 1
            assert page.locator('.device-name').inner_text() == 'Audio console'
            page.locator('#system-buttons').get_by_role('button', name='Video', exact=True).click()
            assert page.locator('#no-results').is_visible()
            page.locator('#venue-filter').select_option('')
            assert page.locator('.device-name').inner_text() == 'ATEM main'
            page.get_by_role('button', name='Clear', exact=True).click()
            assert page.locator('.device-row').count() == 2
            assert page.locator('#system-buttons').get_by_role('button', name='All systems', exact=True).get_attribute('aria-pressed') == 'true'
            # Editing reflects the stored system dropdown and venue button.
            page.get_by_role('button', name='Edit ATEM main').click()
            assert page.locator('#form-system').input_value() == 'Video'
            assert page.locator('#form-venue-buttons').get_by_role('button', name='Liquid Lounge', exact=True).get_attribute('aria-pressed') == 'true'
            page.get_by_role('button', name='Cancel', exact=True).click()
            # Import an existing-style workbook with a cover sheet and custom columns.
            book = Workbook()
            book.active.title = 'Cover'
            book.active.append(['AV list'])
            sheet = book.create_sheet('Network')
            sheet.append(['Equipment list'])
            sheet.append(['Equipment', 'Address on network', 'VLAN ID', 'Room', 'Department'])
            sheet.append(['Excel camera', '10.24.176.68', 1500, 'Liquid Lounge', 'video'])
            path = Path(temp) / 'devices.xlsx'
            book.save(path)
            page.locator('#import-file').set_input_files(str(path))
            page.locator('#sheet-choice').get_by_role('option', name='Network').wait_for(state='attached')
            page.locator('#sheet-choice').select_option('Network')
            page.locator('#header-row').fill('2')
            assert page.locator('#confirm-import').is_disabled()
            page.get_by_role('button', name='Reload preview').click()
            page.locator('#map-ip').get_by_role('option', name='2. Address on network', exact=True).wait_for(state='attached')
            page.locator('#map-ip').select_option('1')
            assert not page.locator('[id^=default-]').count()
            assert 'Excel camera' in page.locator('#spreadsheet-preview').inner_text()
            assert len(app.inventory()) == 2  # Preview never writes.
            page.screenshot(path='/tmp/broadcasthub-excel-mobile.png', full_page=True)
            page.locator('#confirm-import').click()
            accept_rows(1)
            page.get_by_role('button', name='Edit Excel camera').wait_for()
            assert len(app.inventory()) == 3
            assert next(d for d in app.inventory() if d['name'] == 'Excel camera')['category'] == ''
            # A supplied invalid address must leave all existing records intact.
            path = Path(temp) / 'extra.csv'
            path.write_text('Device,IP Address,VLAN\nCSV device,invalid,1500\n')
            page.locator('#import-file').set_input_files(str(path))
            page.locator('#map-venue').wait_for()
            page.locator('#confirm-import').click()
            page.locator('#row-review-dialog').wait_for(state='visible')
            assert page.locator('#row-review-ip-hint').is_visible()
            assert page.locator('#review-ip').input_value() == 'invalid'
            page.locator('#accept-review-row').click()
            page.locator('#row-review-error').wait_for(state='visible')
            assert 'valid IPv4' in page.locator('#row-review-error').inner_text()
            assert len(app.inventory()) == 3
            page.locator('#cancel-row-review').click()
            page.locator('#spreadsheet-dialog').wait_for(state='hidden')
            # IPTV has separate records, source choices, and multicast support.
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            assert page.locator('#total').inner_text() == '0'
            page.get_by_role('button', name='Add channel', exact=True).click()
            page.locator('[name=name]').fill('Ship information')
            page.locator('[name=ip]').fill('239.1.1.10')
            page.locator('[name=port]').fill('1234')
            page.locator('#form-source-buttons').get_by_role('button', name='Onboard', exact=True).click()
            page.get_by_role('button', name='Save channel', exact=True).click()
            page.get_by_role('button', name='Edit Ship information').wait_for()
            assert page.locator('.device-row').count() == 1
            assert page.locator('.ip-confirm').count() == 0
            assert page.locator('#venue-filter-group').is_hidden()
            assert page.locator('#vlan-filter-label').count() == 0
            assert 'Port 1234' in page.locator('.device-row').inner_text()
            page.locator('#source-buttons').get_by_role('button', name='Satellite', exact=True).click()
            assert page.locator('#no-results').is_visible()
            page.get_by_role('button', name='Clear', exact=True).click()
            page.get_by_role('button', name='Add channel', exact=True).click()
            page.locator('[name=name]').fill('BBC News')
            page.locator('[name=ip]').fill('239.1.1.11')
            page.locator('[name=port]').fill('1234')
            page.locator('#form-source-buttons').get_by_role('button', name='Satellite', exact=True).click()
            page.get_by_role('button', name='Save channel', exact=True).click()
            page.get_by_role('button', name='Edit BBC News').wait_for()
            assert page.locator('#system-count').inner_text() == '1 / 1'
            page.get_by_role('button', name='Edit BBC News').click()
            assert page.locator('#form-source-buttons').get_by_role('button', name='Satellite', exact=True).get_attribute('aria-pressed') == 'true'
            assert page.locator('[name=port]').input_value() == '1234'
            assert page.locator('#venue-field').is_hidden()
            assert page.locator('#vlan-field').is_hidden()
            page.get_by_role('button', name='Cancel', exact=True).click()
            path = Path(temp) / 'iptv.csv'
            path.write_text('Name,IP Address,Port,Channel Name,MCAST IP [S],MCAST PORT [S],Source,Category,Inventory type,Notes\n'
                            'Wrong name,invalid,bad,Movie channel,239.1.1.12,5000,Onboard,Ignored category,device,Imported channel note\n'
                            'Wrong name,invalid,bad,Skipped channel,239.1.1.13,5001,Satellite,Ignored category,device,Imported channel note\n'
                            ',,,,,,Invalid source,Ignored category,unknown,\n')
            page.locator('#import-file').set_input_files(str(path))
            page.locator('#map-port').wait_for()
            assert page.locator('#column-mappings select').evaluate_all('(nodes) => nodes.map(n => n.id)') == ['map-name','map-ip','map-port','map-notes']
            assert page.locator('#map-name').input_value() == '3'
            assert page.locator('#map-ip').input_value() == '4'
            assert page.locator('#map-port').input_value() == '5'
            before = len(app.inventory())
            page.locator('#confirm-import').click()
            page.locator('#row-review-dialog').wait_for(state='visible')
            assert page.locator('#row-review-fields input').evaluate_all('(nodes) => nodes.map(n => n.id)') == ['review-name','review-ip','review-port']
            assert page.locator('#row-review-fields > label').evaluate_all('(nodes) => nodes.map(n => n.childNodes[0].textContent)') == ['Channel Name','MCAST IP [S]','MCAST PORT [S]','Notes']
            assert page.locator('#review-ip').input_value() == '239.1.1.12'
            assert page.locator('#review-port').input_value() == '5000'
            assert page.locator('#review-notes').input_value() == 'Imported channel note'
            page.locator('#review-notes').fill('Reviewed channel notes\nSatellite rack')
            assert len(app.inventory()) == before
            page.locator('#accept-review-row').click()
            page.wait_for_function("() => document.getElementById('row-review-progress').textContent.includes('2 of 2')")
            assert len(app.inventory()) == before + 1
            page.locator('#skip-review-row').click()
            page.locator('#spreadsheet-dialog').wait_for(state='hidden')
            page.get_by_role('button', name='Edit Movie channel').wait_for()
            imported = next(d for d in app.inventory() if d['name'] == 'Movie channel')
            assert imported['record_type'] == 'iptv' and imported['ip'] == '239.1.1.12' and imported['port'] == 5000
            assert imported['category'] == imported['channel_source'] == imported['venue'] == ''
            assert imported['notes'] == 'Reviewed channel notes\nSatellite rack'
            assert not any(d['name'] == 'Skipped channel' for d in app.inventory())
            assert not page.locator('.ip-confirm').count()
            assert page.locator('.device-row').count() == 3
            page.screenshot(path='/tmp/broadcasthub-iptv-mobile.png', full_page=True)
            page.get_by_role('tab', name='AV devices', exact=True).click()
            assert page.locator('.device-row').count() == 3
            assert page.get_by_role('button', name='Edit Ship information').count() == 0
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            assert page.get_by_role('button', name='Edit Ship information').is_visible()
            assert page.locator('.ip-confirm').count() == 0
            # General equipment inventory is separate and has no network fields.
            page.get_by_role('tab', name='Equipment inventory', exact=True).click()
            page.locator('#equipment-empty').wait_for(state='visible')
            page.locator('#add-device').click()
            page.locator('#equipment-form [name=brand]').fill('Blackmagic Design')
            page.locator('#equipment-form [name=model]').fill('ATEM Mini Pro')
            page.locator('#equipment-form [name=description]').fill('HDMI switcher')
            page.locator('#equipment-form [name=serial_number]').fill('ATEM-001')
            page.locator('#equipment-form [name=quantity]').fill('1')
            page.locator('#equipment-form [name=location]').fill('Broadcast center')
            page.locator('#equipment-form [name=notes]').fill('With PSU')
            page.get_by_role('button', name='Save equipment', exact=True).click()
            page.get_by_role('button', name='Edit equipment Blackmagic Design ATEM Mini Pro').wait_for()
            assert page.locator('#equipment-units').inner_text() == '1'
            assert page.locator('#equipment-rows tr').count() == 1
            assert page.locator('#inventory').is_hidden()
            page.locator('#equipment-search').fill('missing')
            assert page.locator('#equipment-no-results').is_visible()
            page.locator('#equipment-clear').click()
            page.get_by_role('button', name='Edit equipment Blackmagic Design ATEM Mini Pro').click()
            page.locator('#equipment-form [name=quantity]').fill('2')
            page.locator('#equipment-dialog').get_by_role('button', name='Save changes', exact=True).click()
            page.locator('#equipment-dialog').wait_for(state='hidden')
            page.wait_for_function("() => document.getElementById('equipment-units').textContent === '2'")
            path = Path(temp) / 'equipment.csv'
            path.write_text('BRAND,MODEL,DESCRIPTION,SERIAL NUMBER,QUANTITY,VENUE,NOTES\nNeutrik,XLR,Audio cable,,20,RD Storeroom,Spare stock\n')
            page.locator('#import-file').set_input_files(str(path))
            page.locator('#map-brand').wait_for()
            page.get_by_label('Edit location Storeroom', exact=True).fill('Store 2')
            assert 'Store 2' in page.locator('#spreadsheet-preview').inner_text()
            assert len(app.equipment_inventory()) == 1
            page.locator('#confirm-import').click()
            accept_rows(1)
            page.get_by_role('button', name='Edit equipment Neutrik XLR').wait_for()
            assert page.locator('#equipment-units').inner_text() == '22'
            page.locator('#equipment-location-filter').select_option('Store 2')
            assert page.locator('#equipment-rows tr').count() == 1
            page.locator('#equipment-clear').click()
            open_export_menu('equipment-panel')
            with page.expect_download() as download:
                page.locator('#equipment-panel').get_by_role('link', name='JSON ↓').click()
            equipment_export = json.loads(Path(download.value.path()).read_text())
            assert len(equipment_export['equipment']) == 2
            page.get_by_role('button', name='Delete equipment Neutrik XLR').click()
            page.get_by_role('button', name='Delete equipment', exact=True).click()
            page.wait_for_function("() => document.getElementById('equipment-units').textContent === '2'")
            path = Path(temp) / 'equipment.json'
            path.write_text(json.dumps(equipment_export))
            page.locator('#import-file').set_input_files(str(path))
            page.get_by_role('button', name='Edit equipment Neutrik XLR').wait_for()
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            assert page.locator('.device-row').count() == 3
            page.get_by_role('tab', name='AV devices', exact=True).click()
            assert page.locator('.device-row').count() == 3
            page.get_by_role('tab', name='Equipment inventory', exact=True).click()
            page.get_by_role('button', name='Edit equipment Neutrik XLR').wait_for()
            assert page.locator('#equipment-rows tr').count() == 2
            assert 'avtrack' in page.title()
            # Requested CSV workflow: prioritize exact headers, clean venues,
            # skip repeated IPs across VLANs, and keep devices with text VLANs.
            page.get_by_role('tab', name='AV devices', exact=True).click()
            path = Path(temp) / 'venue-devices.csv'
            path.write_text('Device ID,o.O,Name,Location,IP,DEVICE NAME,VENUE,IP Adress,VLAN,Notes,Category,Inventory type,Channel source,Port,System\n'
                            'old-id,unwanted,Wrong name,Wrong location,invalid,CSV switcher,RD MAIN LOUNGE,10.24.176.90,1500,o.O\n'
                            'old-id,unwanted,Wrong name,Wrong location,invalid,CSV camera,rd MAIN LOUNGE,10.24.176.91,o.O,Imported camera notes,Ignored category,unknown,Cable,bad port,bad system\n'
                            'old-id,unwanted,Wrong name,Wrong location,invalid,CSV lights,RD POOL DECK,10.24.176.92,1.5,o.O\n'
                            'old-id,unwanted,Wrong name,Wrong location,invalid,Repeated switcher,RD MAIN LOUNGE,10.24.176.90,1501,o.O\n'
                            'old-id,unwanted,Wrong name,Wrong location,invalid,Existing ATEM,RD MAIN LOUNGE,10.24.176.66,1502,o.O\n')
            page.locator('#import-file').set_input_files(str(path))
            page.locator('#confirm-import').wait_for(state='visible')
            page.wait_for_function("() => !document.getElementById('confirm-import').disabled")
            assert page.locator('#column-mappings select').evaluate_all('(nodes) => nodes.map(n => n.id)') == ['map-venue','map-name','map-ip','map-vlan','map-notes']
            assert page.locator('#map-venue').input_value() == '4'
            assert page.locator('#map-name').input_value() == '3'
            assert page.locator('#map-ip').input_value() == '5'
            assert 'Ignored columns: Device ID, o.O' in page.locator('#spreadsheet-summary').inner_text()
            assert not page.locator('#map-name option').filter(has_text='Device ID').count()
            assert not page.locator('#map-name option').filter(has_text='o.O').count()
            assert 'RD MAIN' not in page.locator('#spreadsheet-preview').inner_text()
            assert 'Missing' not in page.locator('#spreadsheet-preview').inner_text()
            assert page.locator('#import-venue-list button').all_text_contents() == ['MAIN LOUNGE', 'POOL DECK']
            page.locator('#import-venue-list').get_by_role('button', name='MAIN LOUNGE', exact=True).click()
            venue_input = page.get_by_label('Edit venue MAIN LOUNGE', exact=True)
            assert venue_input.evaluate('(node) => node === document.activeElement')
            venue_input.fill('Main Lounge')
            assert 'Main Lounge' in page.locator('#spreadsheet-preview').inner_text()
            assert 'MAIN LOUNGE' not in page.locator('#spreadsheet-preview').inner_text()
            preview_colors = page.locator('#import-venue-list button').evaluate_all('(nodes) => nodes.map(n => getComputedStyle(n).color)')
            assert len(set(preview_colors)) == 2
            assert len([r for r in app.inventory() if r['record_type'] == 'device']) == 3
            page.screenshot(path='/tmp/broadcasthub-venue-preview-mobile.png', full_page=True)
            page.locator('#confirm-import').click()
            accept_rows(3, total=5)
            page.get_by_role('button', name='Edit CSV camera', exact=True).wait_for()
            assert len([r for r in app.inventory() if r['record_type'] == 'device']) == 6
            assert next(r['notes'] for r in app.inventory() if r['name'] == 'CSV camera') == 'Imported camera notes'
            assert all(r['notes'] == '' for r in app.inventory() if r['name'] in ['CSV switcher', 'CSV lights'])
            assert all(r['category'] == '' for r in app.inventory() if r['name'].startswith('CSV '))
            assert {r['name']:r['discipline'] for r in app.inventory() if r['name'].startswith('CSV ')} == {'CSV switcher':'Network','CSV camera':'Video','CSV lights':'Lighting'}
            venues = page.locator('#venue-filter option').all_text_contents()
            assert venues[-2:] == ['Main Lounge', 'POOL DECK'], venues
            assert venues.count('Main Lounge') == 1
            assert all(r['venue'] == 'Main Lounge' for r in app.inventory() if r['name'] in ['CSV camera', 'CSV switcher'])
            colors = page.locator('#venue-filter option.venue-choice').evaluate_all('(nodes) => nodes.map(n => getComputedStyle(n).color)')
            assert len(set(colors)) == len(colors)
            warnings = page.locator('#import-warnings')
            assert warnings.is_visible()
            assert 'Row 5:' in warnings.inner_text() and '10.24.176.90' in warnings.inner_text()
            assert 'Row 6:' in warnings.inner_text() and '10.24.176.66' in warnings.inner_text()
            assert page.locator('#import-warning-list li').count() == 2
            assert not page.get_by_role('button', name='Edit Repeated switcher', exact=True).count()
            main_venue = page.locator('#venue-filter')
            main_color = main_venue.locator('option[value="Main Lounge"]').evaluate('(n) => getComputedStyle(n).color')
            assert main_color == preview_colors[0]
            main_venue.select_option('Main Lounge')
            assert main_venue.input_value() == 'Main Lounge'
            assert main_venue.evaluate('(n) => getComputedStyle(n).color') == main_color
            assert page.locator('.device-row').count() == 2
            assert page.locator('.device-row').filter(has_text='CSV camera').locator('.ip-cell .cell-caption').inner_text() == ''
            page.get_by_role('button', name='Confirm IP for CSV camera', exact=True).click()
            page.get_by_role('button', name='IP confirmed for CSV camera', exact=True).wait_for()
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            assert not warnings.is_visible()
            page.get_by_role('tab', name='AV devices', exact=True).click()
            assert warnings.is_visible()
            page.get_by_role('button', name='Refresh', exact=True).click()
            page.get_by_role('button', name='Edit CSV lights', exact=True).wait_for()
            assert warnings.is_visible()
            page.get_by_role('button', name='Edit CSV lights', exact=True).click()
            assert page.locator('#form-venue-buttons').get_by_role('button', name='Main Lounge', exact=True).evaluate('(n) => getComputedStyle(n).color') == main_color
            page.locator('[name=ip]').fill('10.24.176.91')
            page.get_by_role('button', name='Save changes', exact=True).click()
            page.locator('#form-error').wait_for(state='visible')
            assert 'Each AV IP must be unique' in page.locator('#form-error').inner_text()
            page.get_by_role('button', name='Cancel', exact=True).click()
            # Missing cells and columns remain blank, with no defaults, and can
            # be completed through Edit without requiring other missing fields.
            path = Path(temp) / 'partial-devices.csv'
            path.write_text('DEVICE NAME,VENUE,IP Adress,Notes\n'
                            'Unassigned camera,,,\n'
                            ',RD BACKSTAGE,,\n'
                            ',,10.24.176.93,\n'
                            ',,,Find this device later\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => document.getElementById('spreadsheet-dialog').open && !document.getElementById('confirm-import').disabled")
            assert not page.locator('[id^=default-]').count()
            assert not page.locator('#map-category, #map-discipline, #map-record_type, #map-channel_source, #map-port').count()
            assert 'Missing' not in page.locator('#spreadsheet-preview').inner_text()
            page.locator('#confirm-import').click()
            accept_rows(4)
            assert any(row['notes'] == 'Find this device later' for row in app.inventory())
            page.get_by_role('button', name='Edit Unassigned camera', exact=True).wait_for()
            unassigned = page.locator('.device-row').filter(has_text='Unassigned camera')
            assert unassigned.locator('.device-ip').inner_text() == ''
            assert unassigned.locator('.device-category').inner_text() == ''
            assert unassigned.locator('.device-venue').inner_text() == ''
            assert not unassigned.locator('.badge, .ip-confirm').count()
            assert not warnings.is_visible()
            system_select = page.get_by_role('combobox', name='System for Unassigned camera', exact=True)
            assert system_select.get_attribute('data-value') == 'Video'
            system_select.click()
            page.locator('.system-options').get_by_role('option', name='Audio', exact=True).click()
            page.wait_for_function("() => document.querySelector('button[aria-label=\"System for Unassigned camera\"]').dataset.value === 'Audio'")
            system_select.click()
            page.locator('.system-options').get_by_role('option', name='Video', exact=True).click()
            page.wait_for_function("() => !document.querySelector('button[aria-label=\"System for Unassigned camera\"]').disabled")
            assert next(r for r in app.inventory() if r['name'] == 'Unassigned camera')['discipline'] == 'Video'
            notes_input = page.get_by_role('textbox', name='Notes for Unassigned camera', exact=True)
            notes_input.fill('Rack B, review later')
            with page.expect_response(lambda response: response.url.endswith('/notes') and response.request.method == 'POST'):
                notes_input.press('Tab')
            assert next(r for r in app.inventory() if r['name'] == 'Unassigned camera')['notes'] == 'Rack B, review later'
            page.get_by_role('button', name='Edit Unassigned camera', exact=True).click()
            assert page.locator('#device-form [name=notes]').input_value() == 'Rack B, review later'
            for field in ['category','venue','ip','vlan']:
                assert page.locator(f'#device-form [name={field}]').input_value() == ''
            assert page.locator('#form-system').input_value() == 'Video'
            page.locator('#device-form [name=ip]').fill('10.24.176.94')
            page.get_by_role('button', name='Save changes', exact=True).click()
            page.get_by_role('button', name='Confirm IP for Unassigned camera', exact=True).wait_for()
            assert next(r for r in app.inventory() if r['name'] == 'Unassigned camera')['category'] == ''
            backstage = next(r for r in app.inventory() if r['venue'] == 'BACKSTAGE')
            page.get_by_role('button', name=f"Edit record {backstage['id']}", exact=True).click()
            page.locator('#device-form [name=name]').fill('Backstage device')
            page.get_by_role('button', name='Save changes', exact=True).click()
            page.get_by_role('button', name='Edit Backstage device', exact=True).wait_for()
            # Empty cells and entirely ignored columns submit successfully with no prompt.
            before = len(app.inventory())
            for filename, content in [('blank.csv', 'VENUE,DEVICE NAME,IP Adress,VLAN\n,,,\n,,,\n'),
                                      ('ignored-only.csv', 'System,Category\ninvalid,Unknown\n')]:
                path = Path(temp) / filename
                path.write_text(content)
                page.locator('#import-file').set_input_files(str(path))
                page.wait_for_function("() => document.getElementById('spreadsheet-dialog').open && !document.getElementById('confirm-import').disabled")
                assert page.locator('#spreadsheet-error').is_hidden()
                page.locator('#confirm-import').click()
                page.locator('#spreadsheet-dialog').wait_for(state='hidden')
                assert len(app.inventory()) == before
            page.get_by_role('tab', name='Equipment inventory', exact=True).click()
            path = Path(temp) / 'partial-equipment.csv'
            path.write_text('Description\nUnidentified spare\nUnidentified spare\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => document.getElementById('spreadsheet-dialog').open && !document.getElementById('confirm-import').disabled")
            assert page.locator('#map-quantity').input_value() == ''
            page.locator('#confirm-import').click()
            accept_rows(2)
            page.wait_for_function("() => document.getElementById('equipment-records').textContent === '4'")
            partial = page.locator('#equipment-rows tr').filter(has_text='Unidentified spare')
            assert partial.count() == 2
            assert partial.locator('[data-label=Quantity]').all_text_contents() == ['', '']
            spare = next(r for r in app.equipment_inventory() if r['description'] == 'Unidentified spare')
            page.get_by_role('button', name=f"Edit equipment record {spare['id']}", exact=True).click()
            assert page.locator('#equipment-quantity').input_value() == ''
            page.locator('#equipment-brand').fill('Sony')
            page.locator('#equipment-quantity').fill('3')
            page.get_by_role('button', name='Save changes', exact=True).click()
            page.get_by_role('button', name='Edit equipment Sony', exact=True).wait_for()
            assert next(r for r in app.equipment_inventory() if r['id'] == spare['id'])['model'] == ''
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            path = Path(temp) / 'partial-iptv.csv'
            path.write_text('Channel\nUnassigned channel\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => document.getElementById('spreadsheet-dialog').open && !document.getElementById('confirm-import').disabled")
            page.locator('#confirm-import').click()
            accept_rows(1)
            page.get_by_role('button', name='Edit Unassigned channel', exact=True).wait_for()
            assert next(r for r in app.inventory() if r['name'] == 'Unassigned channel')['channel_source'] == ''
            page.get_by_role('tab', name='AV devices', exact=True).click()
            # Every row requires a decision. DHCP stays a shared marker, while
            # text can be corrected, cleared, marked DHCP, or skipped.
            before = len(app.inventory())
            path = Path(temp) / 'review.csv'
            path.write_text('VENUE,DEVICE NAME,IP Address,VLAN\n'
                            'RD DHCP LOUNGE,DHCP camera A,dhcp,1500\n'
                            'RD POOL DECK,DHCP camera B,DHCP,1500\n'
                            'RD DHCP LOUNGE,DHCP camera C,assigned automatically,AV network\n'
                            ',Blank address,not assigned,\n'
                            'RD POOL DECK,Skipped camera,10.24.176.95,1500\n'
                            'RD CONTROL ROOM,Corrected camera,10.24.176.96 rack,1500\n'
                            ',,unknown,\n'
                            'RD CONTROL ROOM,Static camera,10.24.176.97,4095\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => !document.getElementById('confirm-import').disabled")
            page.locator('#confirm-import').click()
            page.locator('#row-review-dialog').wait_for(state='visible')
            assert len(app.inventory()) == before
            assert page.locator('#row-review-fields input').evaluate_all('(nodes) => nodes.map(n => n.id)') == ['review-venue','review-name','review-ip','review-vlan']
            assert page.locator('#review-venue').input_value() == 'DHCP LOUNGE'
            assert page.locator('#review-dhcp').get_attribute('aria-pressed') == 'true'
            assert page.locator('#review-ip').input_value() == 'DHCP'
            assert page.locator('#row-review-ip-hint').is_hidden()
            assert page.locator('#row-review-fields > label').evaluate_all('(nodes) => new Set(nodes.map(n => n.getBoundingClientRect().top)).size') == 1
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            page.screenshot(path='/tmp/broadcasthub-row-review-mobile.png')

            def review_action(action):
                progress = page.locator('#row-review-progress').inner_text()
                page.locator(action).click()
                page.wait_for_function("previous => !document.getElementById('spreadsheet-dialog').open || (document.getElementById('row-review-dialog').open && document.getElementById('row-review-progress').textContent !== previous)", arg=progress)

            review_action('#accept-review-row')
            assert len(app.inventory()) == before + 1  # Yes saves immediately.
            assert 'Row 3' in page.locator('#row-review-progress').inner_text()
            review_action('#accept-review-row')
            assert page.locator('#row-review-ip-hint').is_visible()
            assert len(app.inventory()) == before + 2
            page.locator('#review-dhcp').click()
            assert page.locator('#review-ip').input_value() == 'DHCP'
            review_action('#accept-review-row')
            page.locator('#review-clear-ip').click()
            review_action('#accept-review-row')
            assert next(d for d in app.inventory() if d['name'] == 'Blank address')['ip'] == ''
            review_action('#skip-review-row')
            assert not any(d['name'] == 'Skipped camera' for d in app.inventory())
            page.locator('#review-ip').fill('10.24.176.96')
            page.locator('#review-name').fill('Repaired camera')
            page.locator('#review-venue').fill('RD CONTROL ROOM A')
            review_action('#accept-review-row')
            repaired = next(d for d in app.inventory() if d['name'] == 'Repaired camera')
            assert repaired['ip'] == '10.24.176.96' and repaired['venue'] == 'CONTROL ROOM A'
            page.locator('#review-clear-ip').click()  # Entirely blank: no error/write.
            review_action('#accept-review-row')
            page.locator('#accept-review-row').click()
            page.locator('#row-review-error').wait_for(state='visible')
            assert '1 and 4094' in page.locator('#row-review-error').inner_text()
            assert len(app.inventory()) == before + 5
            page.locator('#review-vlan').fill('1500')
            review_action('#accept-review-row')
            assert len(app.inventory()) == before + 6
            assert page.locator('#import-warnings').is_hidden()
            assert page.locator('#address-buttons').get_by_role('button', name='DHCP', exact=True).count() == 1
            page.locator('#address-buttons').get_by_role('button', name='DHCP', exact=True).click()
            assert page.locator('.device-row').count() == 3
            assert not page.locator('.device-row .ip-confirm').count()
            page.locator('#venue-filter').select_option('DHCP LOUNGE')
            assert page.locator('.device-row').count() == 2
            page.locator('#clear-filters').click()
            # Stopping leaves accepted DHCP rows saved and remaining rows untouched.
            path = Path(temp) / 'stop-review.csv'
            path.write_text('DEVICE NAME,IP Address\nAccepted before stop,DHCP\nNever accepted,DHCP\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => !document.getElementById('confirm-import').disabled")
            page.locator('#confirm-import').click()
            page.locator('#row-review-dialog').wait_for(state='visible')
            review_action('#accept-review-row')
            page.locator('#row-review-dialog').press('Escape')
            page.locator('#spreadsheet-dialog').wait_for(state='hidden')
            assert any(d['name'] == 'Accepted before stop' for d in app.inventory())
            assert not any(d['name'] == 'Never accepted' for d in app.inventory())
            page.reload()
            page.locator('#venue-filter option[value="Main Lounge"]').wait_for(state='attached')
            assert page.get_by_role('combobox', name='System for Unassigned camera', exact=True).get_attribute('data-value') == 'Video'
            assert page.locator('#directory-head span:not(.sr-only)').all_text_contents() == ['DEVICE','VENUE','IP','VLAN','SYSTEM','NOTES']
            assert page.locator('.device-row').first.locator(':scope > div').evaluate_all('(nodes) => nodes.map(n => n.className)') == ['device-identity identity-cell','venue-cell','ip-cell','vlan-cell','system-cell','notes-cell','row-actions']
            directory = page.locator('#device-directory')
            header_top = page.locator('#directory-head').evaluate('(node) => node.getBoundingClientRect().top')
            directory.evaluate('(node) => { node.scrollTop = 200; }')
            assert abs(page.locator('#directory-head').evaluate('(node) => node.getBoundingClientRect().top') - header_top) < 1
            directory.evaluate('(node) => { node.scrollTop = 0; node.scrollLeft = 0; }')
            picker = page.get_by_role('combobox', name='System for Unassigned camera', exact=True)
            picker.click()
            choices = page.locator('.system-options [role=option]')
            assert choices.count() == 7
            assert len(set(choices.evaluate_all('(nodes) => nodes.slice(1).map(n => getComputedStyle(n).color)'))) == 6
            assert picker.get_attribute('aria-expanded') == 'true'
            page.locator('.system-options').get_by_role('option', name='Video', exact=True).press('Escape')
            assert picker.get_attribute('aria-expanded') == 'false'
            assert not page.locator('.system-options').count()
            picker.press('ArrowDown')
            page.locator('.system-options').get_by_role('option', name='Video', exact=True).wait_for()
            page.screenshot(path='/tmp/broadcasthub-system-picker-mobile.png')
            page.locator('#iptv-tab').evaluate('(node) => node.click()')
            assert not page.locator('.system-options').count()
            assert page.locator('#directory-venue-title').is_hidden()
            assert page.locator('#directory-vlan-title').is_hidden()
            assert page.locator('#directory-system-title').inner_text() == 'SOURCE'
            page.get_by_role('tab', name='AV devices', exact=True).click()
            brand = page.locator('.site-footer strong')
            page.evaluate('document.fonts.ready')
            assert 'JetBrains Mono' in brand.evaluate('(node) => getComputedStyle(node).fontFamily')
            assert 'Space Grotesk' in page.locator('#hero-title').evaluate('(node) => getComputedStyle(node).fontFamily')
            assert 'DM Sans' in page.locator('body').evaluate('(node) => getComputedStyle(node).fontFamily')
            for family in ['Space Grotesk','DM Sans','JetBrains Mono']:
                assert page.evaluate('name => document.fonts.check(`12px "${name}"`) && Array.from(document.fonts).some(font => font.family.includes(name) && font.status === "loaded")', family)
            assert brand.inner_text() == 'avtrack'
            assert brand.evaluate('(node) => getComputedStyle(node).animationName') == 'none'
            assert page.locator('.topbar .brand-logo').inner_text() == 'avtrack'
            assert 'Space Grotesk' in page.locator('.topbar .brand-logo').evaluate('(node) => getComputedStyle(node).fontFamily')
            assert page.locator('link[rel=icon]').get_attribute('href') == '/favicon.svg'
            assert not page.locator('.topbar .brand-word').count()
            assert page.locator('.topbar .record-o').count() == 0
            assert not page.locator('.site-footer .record-o').count()
            page.locator('.topbar .brand').screenshot(path='/tmp/broadcast-title-rec-alignment.png')
            notes_input = page.get_by_role('textbox', name='Notes for Unassigned camera', exact=True)
            assert notes_input.input_value() == 'Rack B, review later'
            notes_input.fill('')
            with page.expect_response(lambda response: response.url.endswith('/notes') and response.request.method == 'POST'):
                notes_input.press('Tab')
            assert next(r for r in app.inventory() if r['name'] == 'Unassigned camera')['notes'] == ''
            assert page.locator('#venue-filter option[value="Main Lounge"]').evaluate('(n) => getComputedStyle(n).color') == main_color
            page.locator('#address-buttons').get_by_role('button', name='DHCP', exact=True).click()
            assert page.locator('.device-row').count() == 4
            assert not page.locator('.device-row .ip-confirm').count()
            page.locator('#clear-filters').click()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            page.screenshot(path='/tmp/broadcasthub-mobile.png', full_page=True)
            page.set_viewport_size({'width':1440,'height':1000})
            directory.evaluate('(node) => { node.scrollTop = 0; node.scrollLeft = 0; }')
            assert directory.evaluate('(node) => node.scrollWidth <= node.clientWidth')
            assert page.locator('.device-row').first.locator(':scope > div').evaluate_all('(nodes) => new Set(nodes.map(n => n.getBoundingClientRect().top)).size') == 1
            assert page.locator('.ip-info-line').evaluate_all('(nodes) => nodes.filter(n => n.querySelector(".ip-confirm")).every(n => { const ip = n.querySelector(".device-ip").getBoundingClientRect(); const button = n.querySelector(".ip-confirm").getBoundingClientRect(); return Math.abs((ip.top + ip.bottom) / 2 - (button.top + button.bottom) / 2) < 1 && button.left >= ip.right; })')
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            page.screenshot(path='/tmp/broadcasthub-desktop.png', full_page=True)
            path = Path(temp) / 'desktop-review.csv'
            path.write_text('VENUE,DEVICE NAME,IP Address,VLAN\nRD CONTROL ROOM,Desktop review,DHCP,1500\n')
            before = len(app.inventory())
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => !document.getElementById('confirm-import').disabled")
            page.locator('#confirm-import').click()
            page.locator('#row-review-dialog').wait_for(state='visible')
            assert page.locator('#row-review-fields > label').evaluate_all('(nodes) => new Set(nodes.map(n => n.getBoundingClientRect().top)).size') == 1
            assert page.locator('.review-line-scroll').evaluate('(node) => node.scrollWidth <= node.clientWidth')
            page.screenshot(path='/tmp/broadcasthub-row-review-desktop.png')
            page.locator('#skip-review-row').click()
            page.locator('#spreadsheet-dialog').wait_for(state='hidden')
            assert len(app.inventory()) == before
            # Exporting immediately after editing a note waits for that save.
            notes_input = page.get_by_role('textbox', name='Notes for Unassigned camera', exact=True)
            notes_input.fill('Exported immediately after edit')
            open_export_menu('inventory')
            with page.expect_download() as download:
                page.locator('#inventory').get_by_role('link', name='XLSX ↓', exact=True).click()
            assert download.value.suggested_filename == 'avtrack-av-devices.xlsx'
            workbook = load_workbook(io.BytesIO(Path(download.value.path()).read_bytes()))
            assert workbook.sheetnames == ['AV devices']
            exported = next(row for row in list(workbook['AV devices'].values)[1:] if row[1] == 'Unassigned camera')
            assert exported[4] == 'Video' and exported[5] == 'Exported immediately after edit'
            open_export_menu('inventory')
            with page.expect_download() as download:
                page.locator('#inventory').get_by_role('link', name='PDF ↓', exact=True).click()
            assert download.value.suggested_filename == 'avtrack-av-devices.pdf'
            assert Path(download.value.path()).read_bytes().startswith(b'%PDF-')
            page.get_by_role('tab', name='Equipment inventory', exact=True).click()
            for extension, label in [('xlsx','XLSX ↓'), ('pdf','PDF ↓')]:
                open_export_menu('equipment-panel')
                with page.expect_download() as download:
                    page.locator('#equipment-panel').get_by_role('link', name=label, exact=True).click()
                assert download.value.suggested_filename == f'avtrack-equipment.{extension}'
                data = Path(download.value.path()).read_bytes()
                assert data.startswith(b'PK') if extension == 'xlsx' else data.startswith(b'%PDF-')
            # Numeric address ordering, missing values, system groups, and per-tab order.
            fixture = [
                ('Sort zulu', '10.99.0.2', 'Video'),
                ('Sort Alpha', '10.99.0.10', 'Audio'),
                ('Sort beta', '10.99.1.1', 'Lighting'),
                ('Sort delta', 'DHCP', ''),
                ('Sort echo', '', 'Control'),
                ('Sort foxtrot', '200.1.1.1', 'Audio'),
                ('', '10.99.0.3', ''),
            ]
            for name, address, system in fixture:
                response = page.request.post(f'http://127.0.0.1:{server.server_port}/api/devices', data={
                    'record_type':'device', 'name':name, 'ip':address, 'discipline':system,
                    'notes':'Sorting fixture', 'venue':'Sort venue', 'vlan':1500,
                })
                assert response.status == 201, response.text()
            page.get_by_role('tab', name='AV devices', exact=True).click()
            page.locator('#refresh').click()
            page.wait_for_function("() => !document.getElementById('refresh').disabled")
            page.locator('#search').fill('Sorting fixture')
            names = lambda: page.locator('.device-row .device-name').all_text_contents()
            original_order = names()
            expectations = {
                'ip-asc':['Sort zulu', '', 'Sort Alpha', 'Sort beta', 'Sort foxtrot', 'Sort delta', 'Sort echo'],
                'ip-desc':['Sort foxtrot', 'Sort beta', 'Sort Alpha', '', 'Sort zulu', 'Sort delta', 'Sort echo'],
                'name-asc':['Sort Alpha', 'Sort beta', 'Sort delta', 'Sort echo', 'Sort foxtrot', 'Sort zulu', ''],
                'name-desc':['Sort zulu', 'Sort foxtrot', 'Sort echo', 'Sort delta', 'Sort beta', 'Sort Alpha', ''],
                'system-asc':['Sort Alpha', 'Sort foxtrot', 'Sort echo', 'Sort beta', 'Sort zulu', '', 'Sort delta'],
                'system-desc':['Sort zulu', 'Sort beta', 'Sort echo', 'Sort Alpha', 'Sort foxtrot', '', 'Sort delta'],
                '':original_order,
            }
            for order, expected in expectations.items():
                page.locator('#sort-order').select_option(order)
                assert names() == expected, (order, names())
            page.locator('#sort-order').select_option('name-asc')
            page.locator('#system-buttons').get_by_role('button', name='Audio', exact=True).click()
            assert names() == ['Sort Alpha','Sort foxtrot']
            page.locator('#clear-filters').click()
            assert page.locator('#sort-order').input_value() == 'name-asc'
            for name, address, source in [('IPTV sort A','239.99.0.10','Satellite'),('IPTV sort Z','239.99.0.2','Onboard')]:
                response = page.request.post(f'http://127.0.0.1:{server.server_port}/api/devices', data={
                    'record_type':'iptv', 'name':name, 'ip':address, 'channel_source':source,
                })
                assert response.status == 201, response.text()
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            page.locator('#refresh').click()
            page.wait_for_function("() => !document.getElementById('refresh').disabled")
            page.locator('#search').fill('IPTV sort')
            page.locator('#sort-order').select_option('ip-asc')
            assert names() == ['IPTV sort Z','IPTV sort A']
            page.locator('#sort-order').select_option('system-desc')
            assert names() == ['IPTV sort A','IPTV sort Z']
            assert page.locator('#sort-order option[value="system-desc"]').inner_text() == 'Source · Z–A'
            page.get_by_role('tab', name='AV devices', exact=True).click()
            assert page.locator('#sort-order').input_value() == 'name-asc'
            page.locator('#refresh').click()
            page.wait_for_function("() => !document.getElementById('refresh').disabled")
            assert page.locator('#sort-order').input_value() == 'name-asc'
            page.locator('#search').fill('Sorting fixture')
            assert names() == expectations['name-asc']
            assert page.locator('.system-select').evaluate_all("nodes => nodes.every(node => { const text = node.querySelector('span').getBoundingClientRect(), arrow = node.querySelector('svg').getBoundingClientRect(); return Math.abs((text.top+text.bottom-arrow.top-arrow.bottom)/2) < 1; })")
            assert not page.locator('.device-row .device-icon').count()
            assert page.locator('.device-row .vlan-cell').all_text_contents() == ['1500'] * 7
            assert page.locator('.device-row').evaluate_all('nodes => nodes.every(node => node.getBoundingClientRect().height < 75)')
            page.locator('#device-directory').screenshot(path='/tmp/broadcast-compact-sorted-directory.png')
            page.set_viewport_size({'width':390,'height':844})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            assert page.locator('#sort-order').is_visible()
            page.locator('#sort-order').select_option('ip-desc')
            assert names() == expectations['ip-desc']
            # Skip saved/just-accepted IPs before review; notes can be edited or blank.
            page.locator('#clear-filters').click()
            before = len(app.inventory())
            original = next(row for row in app.inventory() if row['ip'] == '10.24.176.66')
            path = Path(temp) / 'duplicate-review-notes.csv'
            path.write_text('VENUE,DEVICE NAME,IP Address,VLAN,Notes\n'
                            'RD CONTROL ROOM,Already saved,10.24.176.66,4095,Do not overwrite\n'
                            'RD CONTROL ROOM,New review camera,10.24.177.1,1500,"Spreadsheet notes\nSecond line"\n'
                            'RD CONTROL ROOM,Repeated new camera,10.24.177.1,4095,Do not overwrite\n'
                            'RD CONTROL ROOM,Edited collision,10.24.177.2,1500,\n'
                            'RD CONTROL ROOM,Next review camera,10.24.177.3,1500,\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => !document.getElementById('confirm-import').disabled")
            page.locator('#confirm-import').click()
            page.locator('#row-review-dialog').wait_for(state='visible')
            assert page.locator('#review-name').input_value() == 'New review camera'
            assert '2 of 5' in page.locator('#row-review-progress').inner_text()
            assert page.locator('#review-notes').input_value() == 'Spreadsheet notes\nSecond line'
            page.locator('#review-notes').fill('Reviewed notes\nRack 3')
            review_action('#accept-review-row')
            assert page.locator('#review-name').input_value() == 'Edited collision'
            assert '4 of 5' in page.locator('#row-review-progress').inner_text()
            page.locator('#review-ip').fill('10.24.177.1')
            page.locator('#review-vlan').fill('4095')
            review_action('#accept-review-row')
            assert page.locator('#review-name').input_value() == 'Next review camera'
            assert page.locator('#review-notes').input_value() == ''
            review_action('#accept-review-row')
            assert len(app.inventory()) == before + 2
            assert next(row['notes'] for row in app.inventory() if row['name'] == 'New review camera') == 'Reviewed notes\nRack 3'
            assert next(row['notes'] for row in app.inventory() if row['name'] == 'Next review camera') == ''
            assert next(row for row in app.inventory() if row['id'] == original['id']) == original
            assert page.locator('#import-warning-list li').count() == 3
            # A file containing only known IPs completes without asking Yes/Skip.
            path.write_text('DEVICE NAME,IP Address,VLAN,Notes\n'
                            'Duplicate original,10.24.176.66,4095,Keep saved notes\n'
                            'Duplicate reviewed,10.24.177.1,4095,Keep saved notes\n')
            page.locator('#import-file').set_input_files(str(path))
            page.wait_for_function("() => !document.getElementById('confirm-import').disabled")
            page.locator('#confirm-import').click()
            page.locator('#spreadsheet-dialog').wait_for(state='hidden')
            assert page.locator('#row-review-dialog').is_hidden()
            assert len(app.inventory()) == before + 2
            assert next(row['notes'] for row in app.inventory() if row['name'] == 'New review camera') == 'Reviewed notes\nRack 3'
            # Current-view exports follow venue/system/search filters and display order.
            page.set_viewport_size({'width':1440,'height':1000})
            page.locator('#clear-filters').click()
            page.locator('#venue-filter').select_option('Sort venue')
            page.locator('#system-buttons').get_by_role('button', name='Audio', exact=True).click()
            page.locator('#sort-order').select_option('ip-asc')
            assert names() == ['Sort Alpha', 'Sort foxtrot']

            def scoped_download(panel, label, scope):
                open_export_menu(panel, scope)
                with page.expect_download() as download:
                    page.locator('#'+panel).get_by_role('link', name=label, exact=True).click()
                extension = {'JSON ↓':'json','CSV ↓':'csv','XLSX ↓':'xlsx','PDF ↓':'pdf'}[label]
                prefix = 'equipment' if panel == 'equipment-panel' else ('iptv-channels' if page.locator('#iptv-tab').get_attribute('aria-selected') == 'true' else 'av-devices')
                assert download.value.suggested_filename == f'avtrack-{prefix}.{extension}'
                return Path(download.value.path()).read_bytes()

            raw = scoped_download('inventory', 'JSON ↓', 'filtered')
            assert [row['name'] for row in json.loads(raw)['devices']] == ['Sort Alpha', 'Sort foxtrot']
            raw = scoped_download('inventory', 'CSV ↓', 'filtered')
            assert [row['name'] for row in csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))] == ['Sort Alpha', 'Sort foxtrot']
            raw = scoped_download('inventory', 'XLSX ↓', 'filtered')
            book = load_workbook(io.BytesIO(raw))
            assert [row[1] for row in list(book['AV devices'].values)[1:]] == ['Sort Alpha', 'Sort foxtrot']
            assert book.sheetnames == ['AV devices']
            raw = scoped_download('inventory', 'PDF ↓', 'filtered')
            assert b'http://10.99.0.10' in raw and b'http://200.1.1.1' in raw
            assert b'http://10.99.0.2' not in raw
            page.get_by_role('checkbox', name='Select Sort Alpha', exact=True).check()
            assert page.locator('#select-visible-devices').evaluate('(node) => node.indeterminate')
            assert page.locator('#network-export-scope').input_value() == 'selected'
            page.locator('#clear-filters').click()
            page.locator('#venue-filter').select_option('CONTROL ROOM')
            assert page.locator('#network-selection-summary').is_hidden()
            page.get_by_role('checkbox', name='Select New review camera', exact=True).check()
            exported = json.loads(scoped_download('inventory', 'JSON ↓', 'selected'))['devices']
            assert {row['name'] for row in exported} == {'New review camera'}
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            assert page.locator('#network-selection-summary').is_hidden()
            page.locator('#search').fill('IPTV sort')
            page.get_by_role('checkbox', name='Select IPTV sort A', exact=True).check()
            exported = json.loads(scoped_download('inventory', 'JSON ↓', 'selected'))['devices']
            assert [row['name'] for row in exported] == ['IPTV sort A']
            assert exported[0]['record_type'] == 'iptv'
            page.get_by_role('tab', name='AV devices', exact=True).click()
            assert page.locator('#network-selection-count').inner_text().startswith('1 selected')
            page.locator('#clear-network-selection').click()
            assert page.locator('#network-export-scope').input_value() == 'filtered'
            page.locator('#clear-filters').click()
            page.locator('#venue-filter').select_option('Sort venue')
            page.locator('#select-visible-devices').check()
            assert page.locator('.device-row .row-select:checked').count() == 7
            page.locator('#select-visible-devices').uncheck()
            assert page.locator('#network-selection-summary').is_hidden()
            page.locator('#search').fill('No matching record for export')
            assert json.loads(scoped_download('inventory', 'JSON ↓', 'filtered'))['devices'] == []
            assert len(json.loads(scoped_download('inventory', 'JSON ↓', 'all'))['devices']) == sum(row['record_type']=='device' for row in app.inventory())
            # Equipment selection contains only records visible in its current location filter.
            for brand, location in [('Export camera', 'Export theatre'), ('Export spare', 'Export store')]:
                response = page.request.post(f'http://127.0.0.1:{server.server_port}/api/equipment', data={'brand':brand, 'location':location})
                assert response.status == 201
            page.get_by_role('tab', name='Equipment inventory', exact=True).click()
            page.wait_for_function("() => !document.getElementById('equipment-refresh').disabled")
            page.locator('#equipment-clear').click()
            page.locator('#equipment-location-filter').select_option('Export theatre')
            assert [row['brand'] for row in json.loads(scoped_download('equipment-panel', 'JSON ↓', 'filtered'))['equipment']] == ['Export camera']
            page.locator('#select-visible-equipment').check()
            page.locator('#equipment-location-filter').select_option('Export store')
            page.get_by_role('checkbox', name='Select equipment Export spare', exact=True).check()
            assert page.locator('#equipment-selection-count').inner_text() == '1 selected'
            exported = json.loads(scoped_download('equipment-panel', 'JSON ↓', 'selected'))['equipment']
            assert {row['brand'] for row in exported} == {'Export spare'}
            page.locator('#clear-equipment-selection').click()
            assert page.locator('#equipment-export-scope').input_value() == 'filtered'
            for width in [390, 320]:
                page.set_viewport_size({'width':width,'height':844})
                open_export_menu('equipment-panel', 'filtered')
                assert page.locator('#equipment-panel .export-links').evaluate('(node) => {const r=node.getBoundingClientRect(); return r.left>=0 && r.right<=innerWidth;}')
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                page.keyboard.press('Escape')
                assert page.get_by_role('checkbox', name='Select equipment Export spare', exact=True).is_visible()
            # Batch deletion previews visible selections, supports cancellation, and
            # removes only the confirmed IDs while keeping other inventories intact.
            page.set_viewport_size({'width':1440,'height':1000})
            assert page.locator('#equipment-rows .row-select').evaluate_all('(nodes) => {const x=document.getElementById("select-visible-equipment").getBoundingClientRect().left; return nodes.every(n=>Math.abs(n.getBoundingClientRect().left-x)<1);}')
            for name, venue in [('Batch delete A','Batch theatre'), ('Batch delete B','Batch theatre'), ('Batch keep C','Batch lounge')]:
                response = page.request.post(f'http://127.0.0.1:{server.server_port}/api/devices', data={'name':name,'venue':venue})
                assert response.status == 201
            page.get_by_role('tab', name='AV devices', exact=True).click()
            page.locator('#refresh').click()
            page.wait_for_function("() => !document.getElementById('refresh').disabled")
            page.locator('#clear-filters').click()
            page.locator('#clear-network-selection').click() if page.locator('#network-selection-summary').is_visible() else None
            assert page.locator('.device-row .row-select').evaluate_all('(nodes) => {const x=document.getElementById("select-visible-devices").getBoundingClientRect().left; return nodes.every(n=>Math.abs(n.getBoundingClientRect().left-x)<1);}')
            page.locator('#venue-filter').select_option('Batch theatre')
            page.get_by_role('checkbox', name='Select Batch delete A', exact=True).check()
            page.get_by_role('checkbox', name='Select Batch delete B', exact=True).check()
            before = app.inventory()
            page.locator('#delete-selected-devices').click()
            assert page.locator('#batch-delete-title').inner_text() == 'Delete 2 devices?'
            assert 'outside the current filters' not in page.locator('#batch-delete-description').inner_text()
            assert page.locator('#batch-delete-list li').count() == 2
            assert 'Batch keep C' not in page.locator('#batch-delete-list').inner_text()
            page.locator('#cancel-batch-delete').click()
            assert app.inventory() == before
            page.locator('#delete-selected-devices').click()
            page.locator('#confirm-batch-delete').click()
            page.locator('#batch-delete-dialog').wait_for(state='hidden')
            page.wait_for_function("() => !document.getElementById('confirm-batch-delete').disabled")
            assert len(app.inventory()) == len(before)-2
            assert not any(row['name'] in ['Batch delete A','Batch delete B'] for row in app.inventory())
            assert any(row['name']=='Batch keep C' for row in app.inventory())
            assert page.locator('#network-selection-summary').is_hidden()
            page.get_by_role('tab', name='IPTV channels', exact=True).click()
            if page.locator('#network-selection-summary').is_visible():
                page.locator('#clear-network-selection').click()
            page.locator('#search').fill('IPTV sort')
            page.locator('#select-visible-devices').check()
            page.locator('#delete-selected-devices').click()
            assert page.locator('#batch-delete-title').inner_text() == 'Delete 2 channels?'
            page.locator('#confirm-batch-delete').click()
            page.locator('#batch-delete-dialog').wait_for(state='hidden')
            page.wait_for_function("() => !document.getElementById('confirm-batch-delete').disabled")
            assert not any(row['name'].startswith('IPTV sort') for row in app.inventory())
            assert any(row['name']=='Batch keep C' for row in app.inventory())
            # Searching for decoder and selecting all exports only visible matches.
            for name in ['Decoder cabin 1', 'Decoder cabin 2', 'Encoder keep']:
                response = page.request.post(f'http://127.0.0.1:{server.server_port}/api/devices', data={'name':name})
                assert response.status == 201
            page.get_by_role('tab', name='AV devices', exact=True).click()
            page.locator('#refresh').click()
            page.wait_for_function("() => !document.getElementById('refresh').disabled")
            page.locator('#clear-filters').click()
            page.get_by_role('checkbox', name='Select Encoder keep', exact=True).check()
            page.locator('#search').fill('decoder')
            assert page.locator('#network-selection-summary').is_hidden()
            matching = names()
            assert {'Decoder cabin 1', 'Decoder cabin 2'} <= set(matching)
            page.locator('#select-visible-devices').check()
            assert page.locator('.device-row .row-select:checked').count()==len(matching)
            assert page.locator('#network-export-scope').input_value()=='selected'
            rows = json.loads(scoped_download('inventory', 'JSON ↓', 'selected'))['devices']
            assert [row['name'] for row in rows]==matching
            book = load_workbook(io.BytesIO(scoped_download('inventory', 'XLSX ↓', 'selected')))
            assert book.sheetnames==['AV devices']
            assert [row[1] for row in list(book.active.values)[1:]]==matching
            raw = scoped_download('inventory', 'CSV ↓', 'selected')
            assert [row['name'] for row in csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))]==matching
            raw = scoped_download('inventory', 'PDF ↓', 'selected')
            if shutil.which('pdftotext'):
                text = subprocess.run(['pdftotext','-','-'],input=raw,stdout=subprocess.PIPE,check=True).stdout.decode()
                assert 'Decoder cabin 1' in text and 'Decoder cabin 2' in text
                assert 'Encoder keep' not in text and 'IPTV channels' not in text
            # Every format, including All in this tab, stays within the active tab.
            for tab, record_type, sheet_name, excluded_heading in [
                ('AV devices','device','AV devices','IPTV channels'),
                ('IPTV channels','iptv','IPTV channels','AV devices'),
            ]:
                page.get_by_role('tab', name=tab, exact=True).click()
                page.locator('#clear-filters').click()
                assert page.locator('#inventory .export-links a[download]').evaluate_all('(links) => links.map(link => new URL(link.href).searchParams.get("record_type"))') == [record_type]*4
                expected = [row for row in app.inventory() if row['record_type']==record_type]
                assert page.locator('#network-export-scope option[value="all"]').text_content().endswith(str(len(expected))), (record_type, len(expected), page.locator('#network-export-scope option[value="all"]').text_content())
                rows = json.loads(scoped_download('inventory', 'JSON ↓', 'all'))['devices']
                assert len(rows)==len(expected) and all(row['record_type']==record_type for row in rows)
                raw = scoped_download('inventory', 'CSV ↓', 'all')
                rows = list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
                assert len(rows)==len(expected) and all(row['record_type']==record_type for row in rows)
                book = load_workbook(io.BytesIO(scoped_download('inventory', 'XLSX ↓', 'all')))
                assert book.sheetnames==[sheet_name] and book.active.max_row==len(expected)+1
                raw = scoped_download('inventory', 'PDF ↓', 'all')
                if shutil.which('pdftotext'):
                    text = subprocess.run(['pdftotext','-','-'],input=raw,stdout=subprocess.PIPE,check=True).stdout.decode()
                    assert sheet_name in text and excluded_heading not in text
            assert not errors, errors
            browser.close()
            print('PASS: mobile create, all filters, validation, edit, export, delete, import, reload persistence; filtered/selected downloads across inventories; desktop/mobile overflow; no JS errors.')
    finally:
        server.shutdown()
        server.server_close()

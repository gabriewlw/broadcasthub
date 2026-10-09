"""Named inventories keep imports, item operations, and reports separate."""
import csv
import io
import shutil
import sqlite3
import subprocess
import unittest

from openpyxl import load_workbook
import app
import test_app
from test_equipment import ITEM


class EquipmentInventoryTests(unittest.TestCase):
    setUpClass = test_app.AppTests.__dict__['setUpClass']
    tearDownClass = test_app.AppTests.__dict__['tearDownClass']
    setUp = test_app.AppTests.setUp
    tearDown = test_app.AppTests.tearDown
    request = test_app.AppTests.request

    def create(self, name):
        status, inventory = self.request('/api/equipment/inventories', 'POST', {'name':name})
        self.assertEqual(status, 201)
        return inventory['id']

    def test_create_rename_and_validate_names(self):
        tvs = self.create('  TVs  ')
        self.assertEqual(self.request('/api/equipment/inventories')[1]['inventories'],
                         [{'id':1,'name':'Equipment inventory'}, {'id':tvs,'name':'TVs'}])
        self.assertEqual(self.request('/api/equipment/inventories', 'POST', {'name':'tvs'})[0], 409)
        for name in ['', '  ', 'x'*121, None, 12]:
            self.assertEqual(self.request('/api/equipment/inventories', 'POST', {'name':name})[0], 400)
        self.assertEqual(self.request(f'/api/equipment/inventories/{tvs}', 'PUT', {'name':'Scalas'})[1]['name'], 'Scalas')
        self.assertEqual(self.request(f'/api/equipment/inventories/{tvs}', 'PUT', {'name':'equipment inventory'})[0], 409)
        self.assertEqual(self.request('/api/equipment/inventories/999', 'PUT', {'name':'Missing'})[0], 404)
        self.assertEqual(app.equipment_inventory(tvs), [])

    def test_imports_and_duplicates_are_scoped_to_inventory(self):
        tvs, scalas = self.create('TVs'), self.create('Scalas')
        stock = dict(ITEM, serial_number='', description='Cable')
        for inventory_id in [1, tvs, scalas]:
            payload = dict(version=1, inventory_id=inventory_id, equipment=[ITEM, stock])
            self.assertEqual(self.request('/api/equipment/import', 'POST', payload)[1], {'added':2,'skipped':0})
            self.assertEqual(self.request('/api/equipment/import', 'POST', payload)[1], {'added':0,'skipped':2})
            self.assertEqual(self.request('/api/equipment', 'POST', dict(ITEM, inventory_id=inventory_id))[0], 409)
            rows = self.request(f'/api/equipment?inventory_id={inventory_id}')[1]['equipment']
            self.assertEqual(len(rows), 2)
            self.assertTrue(all(row['inventory_id'] == inventory_id for row in rows))
        self.assertEqual(len({row['id'] for inventory_id in [1,tvs,scalas] for row in app.equipment_inventory(inventory_id)}), 6)

    def test_item_actions_and_located_status_do_not_cross_inventories(self):
        tvs, spare = self.create('TVs'), self.create('Spare parts')
        first = self.request('/api/equipment', 'POST', dict(ITEM, inventory_id=tvs))[1]
        second = self.request('/api/equipment', 'POST', dict(ITEM, inventory_id=spare))[1]
        path = f"/api/equipment/{first['id']}"
        self.assertEqual(self.request(path+'/confirm', 'POST', dict(first, item_confirmed=True))[0], 200)
        self.assertEqual(app.equipment_inventory(tvs)[0]['item_confirmed'], 1)
        self.assertEqual(app.equipment_inventory(spare)[0]['item_confirmed'], 0)
        self.assertEqual(self.request(path+'/confirm', 'POST', dict(first, inventory_id=spare, item_confirmed=False))[0], 404)
        self.assertEqual(self.request(path, 'PUT', dict(first, inventory_id=spare, location='Wrong'))[0], 404)
        self.assertEqual(self.request(path, 'DELETE', {'inventory_id':spare})[0], 404)
        self.assertEqual(app.equipment_inventory(tvs)[0]['location'], ITEM['location'])
        self.assertEqual(self.request(path, 'PUT', dict(first, notes='New note'))[1]['item_confirmed'], 1)
        self.assertEqual(self.request(path, 'DELETE', {'inventory_id':tvs})[0], 200)
        self.assertEqual(app.equipment_inventory(tvs), [])
        self.assertEqual(app.equipment_inventory(spare)[0]['id'], second['id'])

    def test_every_export_format_uses_selected_inventory(self):
        tvs, spare = self.create('TVs'), self.create('Spare parts')
        first = self.request('/api/equipment', 'POST', dict(ITEM, inventory_id=tvs, description='TV item'))[1]
        other = self.request('/api/equipment', 'POST', dict(ITEM, inventory_id=spare, description='Excluded spare', location='Storeroom'))[1]
        for suffix in ['', '.csv', '.xlsx', '.pdf']:
            for method in ['GET', 'POST']:
                with self.subTest(suffix=suffix, method=method):
                    path = '/api/equipment/export'+suffix
                    payload = {'inventory_id':tvs, 'ids':[first['id'],other['id']]} if method == 'POST' else None
                    if method == 'GET':
                        path += f'?inventory_id={tvs}'
                    status, result = self.request(path, method, payload)
                    self.assertEqual(status, 200)
                    if not suffix:
                        self.assertEqual(result['inventory_name'], 'TVs')
                        self.assertEqual([row['description'] for row in result['equipment']], ['TV item'])
                    elif suffix == '.csv':
                        self.assertEqual([row['description'] for row in csv.DictReader(io.StringIO(result.decode('utf-8-sig')))], ['TV item'])
                    elif suffix == '.xlsx':
                        book = load_workbook(io.BytesIO(result))
                        self.assertEqual(book.sheetnames, ['TVs'])
                        self.assertEqual(book.active['A2'].value, 'TV item')
                        self.assertEqual(book.active.max_row, 2)
                    elif shutil.which('pdftotext'):
                        text = subprocess.run(['pdftotext','-','-'], input=result, capture_output=True, check=True).stdout.decode()
                        self.assertIn('TVs', text)
                        self.assertIn('Location: Broadcast center', text)
                        self.assertNotIn('Excluded spare', text)
        self.request(f'/api/equipment/inventories/{tvs}', 'PUT', {'name':"[TVs] / Lounge: HD "*5})
        self.assertEqual(self.request(f'/api/equipment/export.xlsx?inventory_id={tvs}')[0], 200)

    def test_invalid_or_missing_inventory_never_falls_back_to_default(self):
        self.request('/api/equipment', 'POST', ITEM)
        for inventory_id in [0, -1, True, None, 'abc', 2**64]:
            self.assertEqual(self.request('/api/equipment', 'POST', dict(ITEM, inventory_id=inventory_id))[0], 400)
            self.assertEqual(self.request('/api/equipment/import', 'POST', dict(version=1, inventory_id=inventory_id, equipment=[ITEM]))[0], 400)
            self.assertEqual(self.request('/api/equipment/export', 'POST', {'inventory_id':inventory_id,'ids':[1]})[0], 400)
        for path in ['/api/equipment', '/api/equipment/export', '/api/equipment/export.csv', '/api/equipment/export.xlsx', '/api/equipment/export.pdf']:
            self.assertEqual(self.request(path+'?inventory_id=999')[0], 404)
            self.assertEqual(self.request(path+'?inventory_id=')[0], 400)
        self.assertEqual(self.request('/api/equipment/import', 'POST', dict(version=1, inventory_id=999, equipment=[ITEM]))[0], 404)
        self.assertEqual(len(app.equipment_inventory()), 1)

    def test_migration_preserves_ids_quantities_located_flags_and_timestamps(self):
        with sqlite3.connect(app.DB_PATH) as con:
            con.execute("""CREATE TABLE equipment (id INTEGER PRIMARY KEY, brand TEXT NOT NULL,
                model TEXT NOT NULL, description TEXT NOT NULL DEFAULT '', serial_number TEXT NOT NULL DEFAULT '',
                quantity INTEGER NOT NULL, location TEXT NOT NULL, notes TEXT NOT NULL DEFAULT '',
                item_confirmed INTEGER NOT NULL DEFAULT 0, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""")
            con.execute("INSERT INTO equipment VALUES (42,'Sony','TV','Lounge TV','SER-42',3,'Lounge','Keep note',1,'2025-01-01')")
            con.execute("CREATE UNIQUE INDEX equipment_serial ON equipment(serial_number COLLATE NOCASE) WHERE serial_number != ''")
            con.execute("CREATE UNIQUE INDEX equipment_stock_known ON equipment(brand,model,description,location) WHERE serial_number='' ")
        for _ in range(2):
            row = app.equipment_inventory()[0]
            self.assertEqual((row['id'],row['quantity'],row['item_confirmed'],row['inventory_id'],row['notes'],row['updated_at']), (42,3,1,1,'Keep note','2025-01-01'))
        tvs = self.create('TVs')
        self.assertEqual(self.request('/api/equipment', 'POST', dict(ITEM, serial_number='SER-42', inventory_id=tvs))[0], 201)
        self.assertEqual(len(app.equipment_inventory()), 1)

    def test_overview_counts_all_inventories_and_only_confirmed_fixed_av_ips(self):
        self.assertEqual(self.request('/api/overview')[1], {
            'av_total':0, 'av_validated':0, 'iptv_total':0, 'iptv_onboard':0,
            'iptv_satellite':0, 'inventory_total':0, 'inventory_locations':0, 'inventory_count':1})
        av = self.request('/api/devices', 'POST', dict(name='Confirmed TV', ip='10.20.0.1'))[1]
        self.request(f"/api/devices/{av['id']}/confirm", 'POST', dict(ip=av['ip'], vlan=av['vlan']))
        self.request('/api/devices', 'POST', dict(name='DHCP display', ip='DHCP'))
        self.request('/api/devices', 'POST', dict(name='Unassigned display'))
        for source in ['Onboard', 'Satellite', '']:
            self.request('/api/devices', 'POST', dict(name=source or 'Incomplete channel', record_type='iptv', channel_source=source))
        tvs, spare = self.create('TVs'), self.create('Spare parts')
        self.request('/api/equipment', 'POST', dict(description='TV', location='Lounge', inventory_id=tvs))
        self.request('/api/equipment', 'POST', dict(description='Cable', location='Lounge', inventory_id=spare))
        self.request('/api/equipment', 'POST', dict(description='Spare', location='Store', inventory_id=spare))
        self.request('/api/equipment', 'POST', dict(description='Unknown location', inventory_id=spare))
        self.assertEqual(self.request('/api/overview')[1], {
            'av_total':3, 'av_validated':1, 'iptv_total':3, 'iptv_onboard':1,
            'iptv_satellite':1, 'inventory_total':4, 'inventory_locations':2, 'inventory_count':3})
        self.request(f"/api/devices/{av['id']}", 'PUT', dict(name='Moved TV', ip='10.20.0.2'))
        self.assertEqual(self.request('/api/overview')[1]['av_validated'], 0)

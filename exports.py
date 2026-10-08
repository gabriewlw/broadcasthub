"""Readable PDF reports and Excel workbooks generated from saved inventory."""
import io
import ipaddress
import re
import threading
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SYSTEM_COLORS = {'Video': 'DDD5FA', 'Audio': 'CFEADB', 'Lighting': 'F5E5BD',
                 'Control': 'D0E0FF', 'Network': 'C5EFF0', 'Other': 'EAD8E5'}
_font_lock = threading.Lock()


def confirmation(row):
    if row['record_type'] == 'iptv' or row['ip'] in ('', 'DHCP'):
        return ''
    return 'Confirmed' if row['ip_confirmed'] else 'Pending'


def clean_text(value):
    # Excel and PDF do not support these control characters.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))


def workbook(sections):
    book = Workbook()
    book.remove(book.active)
    for title, columns, rows in sections:
        sheet = book.create_sheet(title)
        sheet.append([label for label, field, width in columns])
        for record in rows:
            sheet.append([record.get(field) for label, field, width in columns])
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.value = clean_text(cell.value)
                    # Literal strings must never become executable formulas.
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == 1:
                    cell.font = Font(name='Consolas', bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='C92F38')
                elif sheet.cell(1, cell.column).value == 'System' and cell.value in SYSTEM_COLORS:
                    cell.fill = PatternFill('solid', fgColor=SYSTEM_COLORS[cell.value])
                elif sheet.cell(1, cell.column).value == 'IP confirmation' and cell.value:
                    cell.fill = PatternFill('solid', fgColor='D1EAD7' if cell.value == 'Confirmed' else 'F9E8B2')
        for index, (_, _, width) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = True
        sheet.print_title_rows = '1:1'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def network_xlsx(rows, record_type=None):
    av = [dict(row, confirmation=confirmation(row)) for row in rows if row['record_type'] == 'device']
    channels = [row for row in rows if row['record_type'] == 'iptv']
    sections = [
        ('AV devices', [('Venue','venue',24), ('Device name','name',30), ('IP Address','ip',20),
                        ('VLAN','vlan',10), ('System','discipline',16), ('Notes','notes',55),
                        ('IP confirmation','confirmation',20), ('Category','category',24)], av),
        ('IPTV channels', [('Channel name','name',30), ('IP Address','ip',20), ('Port','port',10),
                           ('Channel source','channel_source',20), ('Notes','notes',55),
                           ('Category','category',24)], channels)]
    if record_type:
        sections = [sections[0 if record_type == 'device' else 1]]
    return workbook(sections)


def equipment_xlsx(rows):
    return workbook([('Equipment', [('Brand','brand',24), ('Model','model',24),
                                   ('Description','description',40), ('Serial number','serial_number',24),
                                   ('Quantity','quantity',12), ('Location','location',26),
                                   ('Notes','notes',55)], rows)])


def pdf_report(title, sections):
    try:
        from PIL import Image as PILImage
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Flowable
    except ImportError:
        raise ValueError('Install PDF dependencies: py -m pip install -r requirements.txt') from None
    # Embed the same locally bundled body, heading, and label fonts as the website.
    with _font_lock:
        if 'HubText' not in pdfmetrics.getRegisteredFontNames():
            fonts = Path(__file__).parent / 'static' / 'fonts'
            pdfmetrics.registerFont(TTFont('HubText', str(fonts / 'DMSans-Regular.ttf')))
            pdfmetrics.registerFont(TTFont('HubBold', str(fonts / 'DMSans-Bold.ttf')))
            pdfmetrics.registerFontFamily('HubText', normal='HubText', bold='HubBold')
            pdfmetrics.registerFont(TTFont('HubDisplay', str(fonts / 'SpaceGrotesk-Bold.ttf')))
            pdfmetrics.registerFont(TTFont('HubMono', str(fonts / 'JetBrainsMono-Regular.ttf')))
            pdfmetrics.registerFont(TTFont('HubMonoMedium', str(fonts / 'JetBrainsMono-Medium.ttf')))
    logo_path = Path(__file__).parent / 'static' / 'brand-logo.png'
    logo = ImageReader(str(logo_path))
    logo_width, logo_height = logo.getSize()
    with PILImage.open(logo_path) as logo_image:
        # Ignore faint transparent padding when aligning the visible lettering.
        logo_bounds = logo_image.convert('RGBA').getchannel('A').point(lambda alpha: 255 if alpha >= 64 else 0).getbbox()
    logo_bounds = logo_bounds or (0, 0, logo_width, logo_height)
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=26, rightMargin=26,
                                 topMargin=18, bottomMargin=64, title=f'Broadcast Hub - {title}',
                                 author='Broadcast Hub')
    body = ParagraphStyle('HubBody', fontName='HubText', fontSize=8, leading=11, wordWrap='CJK', textColor=colors.HexColor('#EEEEE9'))
    header = ParagraphStyle('HubHeader', parent=body, fontName='HubMonoMedium', fontSize=7, leading=9,
                            textColor=colors.HexColor('#C4C7CF'))
    report_info = ParagraphStyle('HubInfo', parent=body, fontName='HubMono', fontSize=7, leading=9)
    heading = ParagraphStyle('HubHeading', fontName='HubDisplay', fontSize=18, leading=24,
                             textColor=colors.HexColor('#EEEEE9'))
    section_style = ParagraphStyle('HubSection', parent=heading, fontSize=9, leading=12, spaceBefore=6, spaceAfter=4, textColor=colors.HexColor('#FF7278'))

    def paragraph(value, style=body):
        return Paragraph(escape(clean_text(value if value is not None else '')).replace('\n', '<br/>'), style)

    tag_colors = {'Video':('#211E2B','#C2B8F6'), 'Audio':('#182720','#9ED3B8'),
                  'Lighting':('#292317','#E5C587'), 'Control':('#1B2539','#96BCFF'),
                  'Network':('#162A2D','#85DCE0'), 'Other':('#2B212D','#D2B4CB')}

    def report_cell(label, value):
        if label == 'SYSTEM' and value in tag_colors:
            return paragraph(value, ParagraphStyle('HubTag', parent=body, textColor=colors.HexColor(tag_colors[value][1])))
        if label == 'CONFIRMATION' and value in ('Confirmed', 'Pending'):
            confirmed = value == 'Confirmed'
            return paragraph(value, ParagraphStyle('HubConfirmation', parent=body,
                fontName='HubBold' if confirmed else 'HubText',
                textColor=colors.HexColor('#F0FFF5' if confirmed else '#E5C587')))
        if label == 'IP' and value not in ('', 'DHCP', None):
            try:
                address = str(ipaddress.IPv4Address(value))
            except (ValueError, TypeError):
                return paragraph(value)
            # Standard external PDF links; the viewer chooses its tab/window behavior.
            return Paragraph(f'<link href="http://{address}" color="#FF7278"><u>{address}</u></link>', body)
        return paragraph(value)

    class BrandHeading(Flowable):
        def __init__(self):
            super().__init__()
            self.width, self.height = 140, 22

        def draw(self):
            canvas = self.canv
            left, top, right, bottom = logo_bounds
            scale = self.width / (right - left)
            canvas.saveState()
            clip = canvas.beginPath()
            clip.rect(0, 4, self.width, (bottom - top) * scale)
            canvas.clipPath(clip, stroke=0, fill=0)
            canvas.drawImage(logo, -left * scale, 4 - (logo_height - bottom) * scale,
                             width=logo_width * scale, height=logo_height * scale, mask='auto')
            canvas.restoreState()

    story = [BrandHeading(), paragraph(title, report_info), Spacer(1, 3)]
    for name, labels, widths, records in sections:
        story.append(paragraph(f'{name} · {len(records)} records', section_style))
        table_rows = [[paragraph(label, header) for label in labels]]
        table_rows += [[report_cell(label, value) for label, value in zip(labels, row)] for row in records]
        if not records:
            table_rows.append([paragraph('No records')] + [paragraph('')] * (len(labels) - 1))
        table = LongTable(table_rows, colWidths=[document.width * width for width in widths],
                          repeatRows=1, splitInRow=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1C2027')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#14171B'), colors.HexColor('#171A20')]),
            ('GRID', (0,0), (-1,-1), .35, colors.HexColor('#373B44')),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING', (0,0), (-1,-1), 7), ('RIGHTPADDING', (0,0), (-1,-1), 7),
            ('TOPPADDING', (0,0), (-1,-1), 7), ('BOTTOMPADDING', (0,0), (-1,-1), 7),
            ('TOPPADDING', (0,0), (-1,0), 4), ('BOTTOMPADDING', (0,0), (-1,0), 4)]))
        if 'SYSTEM' in labels:
            column = labels.index('SYSTEM')
            for index, record in enumerate(records, 1):
                if record[column] in SYSTEM_COLORS:
                    # A tinted chip matches the colored selection in the website.
                    background, _ = tag_colors[record[column]]
                    table.setStyle(TableStyle([('BACKGROUND', (column,index), (column,index), colors.HexColor(background))]))
        if 'CONFIRMATION' in labels:
            column = labels.index('CONFIRMATION')
            for index, record in enumerate(records, 1):
                if record[column] == 'Confirmed':
                    table.setStyle(TableStyle([('BACKGROUND', (column,index), (column,index), colors.HexColor('#216E43'))]))
        story.append(table)

    def footer(canvas, doc):
        canvas.saveState()
        width, height = landscape(A4)
        canvas.setFillColor(colors.HexColor('#0B0D0F'))
        canvas.rect(0, 0, width, height, stroke=0, fill=1)
        canvas.setStrokeColor(colors.HexColor('#272A30'))
        canvas.line(26, 46, width - 26, 46)
        canvas.setFillColor(colors.HexColor('#EF4444'))
        canvas.circle(30, 29, 3, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor('#EEEEE9'))
        canvas.setFont('HubMonoMedium', 9)
        canvas.drawString(40, 26, 'REPORT GENERATED BY BROADCASTHUB')
        canvas.setFillColor(colors.HexColor('#FF7278'))
        canvas.setFont('HubMonoMedium', 9)
        canvas.drawRightString(width - 26, 26, 'broadcastgab.com')
        canvas.linkURL('https://broadcastgab.com', (width - 135, 23, width - 26, 38), relative=0)
        canvas.setFillColor(colors.HexColor('#A0A2AC'))
        canvas.setFont('HubMono', 7)
        canvas.drawString(40, 13, 'Broadcast operations · Saved inventory')
        canvas.drawRightString(width - 26, 13, f'PAGE {doc.page}')
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def network_pdf(rows, record_type=None):
    av, channels = [], []
    for row in rows:
        name = '\n'.join(filter(None, [row['name'], row['category']]))
        if row['record_type'] == 'device':
            av.append([row['venue'], name, row['ip'], row['vlan'], confirmation(row), row['discipline'], row['notes']])
        else:
            channels.append([name, row['ip'], row['port'], row['channel_source'], row['notes']])
    sections = [
        ('AV devices', ['VENUE','DEVICE','IP','VLAN','CONFIRMATION','SYSTEM','NOTES'], [.14,.16,.15,.07,.12,.11,.25], av),
        ('IPTV channels', ['CHANNEL','IP','PORT','SOURCE','NOTES'], [.21,.20,.08,.15,.36], channels)]
    if record_type:
        sections = [sections[0 if record_type == 'device' else 1]]
    title = sections[0][0] if record_type else 'AV devices and IPTV channels'
    return pdf_report(title, sections)


def equipment_pdf(rows):
    fields = ['brand','model','description','serial_number','quantity','location','notes']
    return pdf_report('Equipment inventory', [('Equipment',
        ['BRAND','MODEL','DESCRIPTION','SERIAL NUMBER','QUANTITY','LOCATION','NOTES'],
        [.13,.13,.22,.15,.07,.13,.17], [[row[field] for field in fields] for row in rows])])

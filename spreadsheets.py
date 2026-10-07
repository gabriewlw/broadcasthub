"""Bounded spreadsheet preview; never evaluates formulas or saves uploaded files."""
import base64
import binascii
import csv
import io
import re
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def ignored_header(value):
    normalized = re.sub(r'\s+', ' ', value.lower().replace('_', ' ').replace('-', ' ')).strip()
    return normalized in ('device id', 'o.o', 'o.o value')


def preview(payload):
    if not isinstance(payload, dict):
        raise ValueError('Provide a spreadsheet file.')
    filename = payload.get('filename', '')
    if not isinstance(filename, str) or Path(filename).suffix.lower() not in ('.xlsx', '.csv'):
        raise ValueError('Use .xlsx or .csv. Save older .xls files as .xlsx in Excel first.')
    try:
        raw = base64.b64decode(payload.get('content', ''), validate=True)
    except (binascii.Error, TypeError, ValueError):
        raise ValueError('Invalid file content.') from None
    if len(raw) > 5_000_000 or (not raw and filename.lower().endswith('.xlsx')):
        raise ValueError('Choose a spreadsheet between 1 byte and 5 MB.')
    header_row = payload.get('header_row', 1)
    if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= 50:
        raise ValueError('Header row must be a number from 1 to 50.')
    book = None
    sheets, sheet = [], ''
    try:
        if filename.lower().endswith('.xlsx'):
            with ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 500 or sum(f.file_size for f in archive.infolist()) > 30_000_000:
                    raise ValueError('This workbook is too large when expanded. Export a smaller sheet as CSV.')
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            sheets = book.sheetnames
            sheet = payload.get('sheet') or sheets[0]
            if sheet not in sheets:
                raise ValueError('Choose a worksheet from this workbook.')
            worksheet = book[sheet]
            worksheet.reset_dimensions()
            source = worksheet.iter_rows(values_only=True)
        else:
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                raise ValueError('Save your CSV as CSV UTF-8 in Excel, then try again.') from None
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
            except csv.Error:
                dialect = csv.excel
            source = csv.reader(io.StringIO(text), dialect)
        headers, rows, row_numbers, characters = [], [], [], 0
        kept_columns, ignored_columns, column_count = [], [], None
        for number, raw_row in enumerate(source, 1):
            if number < header_row:
                continue
            # Restrict raw rows too, so millions of empty rows cannot consume unlimited work.
            if number > 20050:
                raise ValueError('Too many spreadsheet rows. Remove unused rows and try again.')
            values = []
            for value in raw_row:
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                value = '' if value is None else str(value).strip()
                if len(value) > 2000:
                    raise ValueError(f'Row {number} contains a cell longer than 2,000 characters.')
                values.append(value)
            while values and not values[-1]:
                values.pop()
            if len(values) > 50:
                raise ValueError('Import supports up to 50 columns. Export only the device columns.')
            characters += sum(map(len, values))
            if characters > 2_000_000:
                raise ValueError('Spreadsheet text is too large. Split the file into smaller imports.')
            if column_count is None:
                if not values:
                    continue
                column_count = len(values)
                kept_columns = [i for i, value in enumerate(values) if not ignored_header(value)]
                ignored_columns = [value for value in values if ignored_header(value)]
                headers = [values[i] or f'Column {i + 1}' for i in kept_columns]
            else:
                if len(values) > column_count:
                    raise ValueError(f'Row {number} has more columns than the chosen header row.')
                values += [''] * (column_count - len(values))
                values = ['' if values[i].lower() == 'o.o' else values[i] for i in kept_columns]
                # Rows containing only IDs or placeholders do not describe a device.
                if not any(values):
                    continue
                if len(rows) >= 10000:
                    raise ValueError('Import supports up to 10,000 devices per file.')
                rows.append(values)
                row_numbers.append(number)
        return {'headers': headers, 'rows': rows, 'row_numbers': row_numbers, 'sheets': sheets, 'sheet': sheet, 'ignored_columns': ignored_columns}
    except (BadZipFile, InvalidFileException, ParseError, DefusedXmlException, KeyError, csv.Error, OSError):
        raise ValueError('Could not read this spreadsheet. Save a fresh .xlsx or CSV UTF-8 copy in Excel.') from None
    finally:
        if book is not None:
            book.close()
